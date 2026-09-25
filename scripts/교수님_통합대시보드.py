# -*- coding: utf-8 -*-
"""교수님용 통합 대시보드 엑셀 (2026-09-23)

윈도우 엑셀 기준. 홈 → 9개 화면(문제·혜택 배치·운영 시간·버스·월영교·기대효과·사후 검증·전국 진단·서식4) + 방문객 설문 + 데이터 출처.
화면마다 같은 틀: 제목 띠 → 이동 메뉴 → 노란 조절 칸(2~5개) → 핵심 숫자 카드 → 차트 → 한 줄 해석·출처.
원자료·계산 시트(d_…)는 숨김. 매크로 없음(목록 선택 + 수식 + 차트).
입력: 앞선 분석 스크립트들의 결과(보고서/교수님_대시보드/data, 성과도출, 전국진단, 교수브리핑, 조사/01·04 등)
출력: 보고서/교수님_대시보드/안동이어드림_데이터대시보드.xlsx
"""
import json, re, html as H
from pathlib import Path
import numpy as np, pandas as pd
import statsmodels.api as sm, statsmodels.formula.api as smf
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.chart import BarChart, LineChart, ScatterChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as CL
from openpyxl.chart.title import Title
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties, RichTextProperties, RegularTextRun, Font as DFont

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / '보고서/교수님_대시보드'
DD = OUT / 'data'
Q = json.loads((ROOT / '보고서/교수브리핑_20260922/수치.json').read_text(encoding='utf-8'))
P = json.loads((ROOT / '보고서/성과도출_20260922/파라미터추정.json').read_text(encoding='utf-8'))
SIM = json.loads((ROOT / '보고서/성과도출_20260922/시뮬레이션결과.json').read_text(encoding='utf-8'))
SV = json.loads((ROOT / '보고서/성과도출_20260922/사후검증_결과.json').read_text(encoding='utf-8'))
PW = json.loads((ROOT / '보고서/성과도출_20260922/순차도입_검정력.json').read_text(encoding='utf-8'))
S3 = json.loads((ROOT / '보고서/성과도출_20260922/성과도출_수치.json').read_text(encoding='utf-8'))
SVY = json.loads((ROOT / 'data/설문/설문_집계.json').read_text(encoding='utf-8'))          # 팀 온라인 설문(scripts/설문_분석.py)

# ═════════════ 스타일 ═════════════
FN = '맑은 고딕'
NAVY, RED, GRAY, BLUE, LIGHT, INPUT, INK = '1F3864', 'D64541', 'A6A6A6', '2E75B6', 'F2F2F2', 'FFE699', '262626'
thin = Side(style='thin', color='BFBFBF')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
F = lambda size=10, bold=False, color=INK, italic=False: Font(name=FN, size=size, bold=bold, color=color, italic=italic)
FILL = lambda c: PatternFill('solid', fgColor=c)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)
TOPL = Alignment(horizontal='left', vertical='top', wrap_text=True)

SHEETS = [('홈', '홈'), ('1_문제', '문제'), ('2_혜택배치', '혜택 배치'), ('3_운영시간', '운영 시간'), ('4_버스', '버스(BIS)'),
          ('5_월영교', '월영교 야간'), ('6_기대효과', '기대효과'), ('7_사후검증', '사후 검증'), ('8_전국진단', '전국 진단'),
          ('9_서식4', '서식4'), ('10_설문', '방문객 설문'), ('11_출처', '데이터 출처')]
q_ = lambda name: f"'{name}'"                                   # 시트 이름 따옴표

wb = Workbook()
wb.remove(wb.active)
W = {name: wb.create_sheet(name) for name, _ in SHEETS}
TAB = {'홈': NAVY, '1_문제': 'C00000', '2_혜택배치': 'C00000', '3_운영시간': 'ED7D31', '4_버스': 'ED7D31', '5_월영교': 'ED7D31',
       '6_기대효과': '548235', '7_사후검증': '548235', '8_전국진단': '7030A0', '9_서식4': '595959', '10_설문': 'ED7D31', '11_출처': '595959'}


def data_sheet(name):
    ws = wb.create_sheet(name)
    ws.sheet_state = 'hidden'
    return ws


def setup(ws, title, message):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.sheet_properties.tabColor = TAB[ws.title]
    ws.column_dimensions['A'].width = 2
    for c in range(2, 21):
        ws.column_dimensions[CL(c)].width = 10.5
    ws.merge_cells('B1:S2')
    c = ws['B1']; c.value = title; c.font = F(18, True, 'FFFFFF'); c.fill = FILL(NAVY); c.alignment = Alignment(vertical='center', indent=1)
    for r in (1, 2):
        for col in range(2, 20):
            ws.cell(r, col).fill = FILL(NAVY)
    ws.row_dimensions[1].height = 22; ws.row_dimensions[2].height = 22
    ws.merge_cells('B3:S3')
    c = ws['B3']; c.value = message; c.font = F(11, True, NAVY); c.alignment = Alignment(vertical='center', indent=1, wrap_text=True)
    ws.row_dimensions[3].height = 34
    # 이동 메뉴 (12칸: B~M)
    for i, (name, lab) in enumerate(SHEETS):
        cell = ws.cell(5, 2 + i, lab)
        cell.hyperlink = f"#{q_(name)}!A1"
        cur = name == ws.title
        cell.font = F(9, True, 'FFFFFF' if cur else NAVY)
        cell.fill = FILL(NAVY if cur else 'DDEBF7')
        cell.alignment = CENTER; cell.border = BOX
    ws.row_dimensions[5].height = 24
    ws.freeze_panes = 'A6'
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.print_area = 'A1:T62'
    ws.page_margins.left = ws.page_margins.right = 0.3; ws.page_margins.top = ws.page_margins.bottom = 0.4


def label(ws, ref, text, size=9.5, bold=True, color='595959'):
    c = ws[ref]; c.value = text; c.font = F(size, bold, color); c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=False)
    return c


def control(ws, row, col, lab, value, fmt=None, options=None, width=2):
    """노란 조절 칸: 위 줄(row-1)에 라벨(오른쪽 빈칸으로 이어서 표시), 아래 줄(row)에 입력(col ~ col+width-1 병합).
    칸 사이를 한 칸씩 띄워 배치한다(2·5·8·11·14열 + width 2)."""
    ws.row_dimensions[row - 1].height = 18; ws.row_dimensions[row].height = 26
    lc = ws.cell(row - 1, col, lab); lc.font = F(9.5, True, NAVY); lc.alignment = Alignment(horizontal='left', vertical='bottom', wrap_text=False)
    width = max(width - 1, 0)
    if width: ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + width)
    vc = ws.cell(row, col, value)
    vc.font = F(11, True, '000000'); vc.fill = FILL(INPUT); vc.alignment = CENTER
    med = Side('medium', color='BF9000')
    for cc in range(col, col + width + 1):
        ws.cell(row, cc).border = Border(top=med, bottom=med, left=med if cc == col else None, right=med if cc == col + width else None)
    if fmt: vc.number_format = fmt
    if options:
        dv = DataValidation(type='list', formula1=options, allow_blank=False, showDropDown=False)
        dv.error = '목록에서 고르세요'; ws.add_data_validation(dv); dv.add(vc.coordinate)
    return f"{q_(ws.title)}!${CL(col)}${row}", vc.coordinate


def kpi(ws, row, col, lab, formula, fmt, sub=None, width=4, color=NAVY):
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + width - 1)
    ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 2, end_column=col + width - 1)
    ws.merge_cells(start_row=row + 3, start_column=col, end_row=row + 3, end_column=col + width - 1)
    a = ws.cell(row, col, lab); a.font = F(9.5, True, '595959'); a.alignment = CENTER
    b = ws.cell(row + 1, col, formula); b.font = F(20, True, color); b.alignment = CENTER; b.number_format = fmt
    s = ws.cell(row + 3, col, sub); s.font = F(8.5, False, '7F7F7F'); s.alignment = CENTER
    for r in range(row, row + 4):
        for c in range(col, col + width):
            ws.cell(r, c).fill = FILL(LIGHT)
    top = Side('medium', color=color)
    for c in range(col, col + width):
        ws.cell(row, c).border = Border(top=top)


def note(ws, row, text, height=None, col1=2, col2=19, size=9, color='595959'):
    ws.merge_cells(start_row=row, start_column=col1, end_row=row, end_column=col2)
    c = ws.cell(row, col1, text); c.font = F(size, False, color); c.alignment = TOPL
    if height: ws.row_dimensions[row].height = height


def _cp(sz, b=False):
    return CharacterProperties(sz=sz, b=b, latin=DFont(typeface=FN), ea=DFont(typeface=FN))


def _rich(sz, b=False):
    return RichText(bodyPr=RichTextProperties(), p=[Paragraph(pPr=ParagraphProperties(defRPr=_cp(sz, b)), endParaRPr=_cp(sz, b))])


def _title(text, sz, rot=None):
    cp = _cp(sz, True)
    body = RichTextProperties(rot=rot, vert='horz') if rot is not None else RichTextProperties()
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])
    return Title(tx=Text(rich=RichText(bodyPr=body, p=[para])), overlay=False)


def _ttext(t):
    return ''.join(r.t for para in t.tx.rich.p for r in (para.r or []))


def style_chart(ch, title, w=17, h=8.5, legend='b'):
    """엑셀에서 축 라벨이 사라지지 않게(delete=False) 하고, 제목·범례가 그림을 덮지 않게 한다."""
    ch.title = _title(title, 1000); ch.width = w; ch.height = h; ch.visible_cells_only = False
    horiz_bar = isinstance(ch, BarChart) and ch.type == 'bar'
    for ax, vertical in ((ch.x_axis, horiz_bar), (ch.y_axis, not horiz_bar)):
        if ax.delete is None: ax.delete = False
        ax.txPr = _rich(800)
        if ax.majorGridlines is not None:
            ax.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill='D9D9D9', w=6350))
        if ax.title is not None:
            ax.title = _title(_ttext(ax.title), 850, rot=-5400000 if vertical else None)
    if legend is None: ch.legend = None
    else:
        ch.legend.position = legend; ch.legend.overlay = False; ch.legend.txPr = _rich(850)
    if isinstance(ch, BarChart):
        for sr in ch.series: sr.invertIfNegative = False
    return ch


def color_series(s, hexc, line=False):
    if line:
        s.graphicalProperties.line.solidFill = hexc; s.graphicalProperties.line.width = 22000; s.smooth = False
    else:
        s.graphicalProperties.solidFill = hexc; s.graphicalProperties.line.solidFill = hexc


def put_table(ws, r0, c0, df, fmts=None, header_fill='D9E1F2', font_size=9):
    for j, col in enumerate(df.columns):
        c = ws.cell(r0, c0 + j, col); c.font = F(font_size, True); c.fill = FILL(header_fill); c.alignment = CENTER; c.border = BOX
    for i, row in enumerate(df.itertuples(index=False), 1):
        for j, v in enumerate(row):
            v = None if (isinstance(v, float) and np.isnan(v)) else v
            c = ws.cell(r0 + i, c0 + j, v); c.font = F(font_size); c.border = BOX; c.alignment = LEFT if isinstance(v, str) else CENTER
            if fmts and j in fmts: c.number_format = fmts[j]
    return r0 + len(df)


def write_df(ws, df, r0=1, c0=1):
    for j, col in enumerate(df.columns):
        ws.cell(r0, c0 + j, col).font = F(9, True)
    for i, row in enumerate(df.itertuples(index=False), 1):
        for j, v in enumerate(row):
            v = None if (isinstance(v, float) and np.isnan(v)) else (v.item() if hasattr(v, 'item') else v)
            ws.cell(r0 + i, c0 + j, v)
    return r0 + len(df)


# ═════════════ 데이터 시트 ═════════════
# d_시군 (문제)
sg = pd.read_csv(DD / '시군_방문소비.csv')
sg = sg.sort_values('체험문화_변화율').reset_index(drop=True)
dS = data_sheet('d_시군'); nS = write_df(dS, sg)
up = pd.read_csv(DD / '안동_업종.csv')                              # 안동 업종(소비 순)·전국 중앙 변화율
su = pd.read_csv(DD / '시군_업종.csv')                              # scripts/교수님_대시보드_데이터.py
su['키'] = su['시군'] + '|' + su['업종']
dUa = data_sheet('d_업종전체'); write_df(dUa, su[['키', '2024', '2026']]); nUa = len(su) + 1
dU = data_sheet('d_업종')

# d_혜택 (읍면동) · d_업체
hb = pd.DataFrame(Q['혜택배치'])
dH = data_sheet('d_혜택'); write_df(dH, hb[['행정동', '방문점유율', '혜택업체']])
bz = pd.DataFrame(P['A_업체']).sort_values('이용', ascending=False)

# d_역 (안동역 1~8월, 연 × 구분 × 열차 × 승하차 × 시간대)
st = pd.read_csv(ROOT / 'data/external/팀원취합_정제/안동역_승하차_월별_열차종류별_long.csv')
st['연'] = st['연월'].str[:4].astype(int); st['월'] = st['연월'].str[5:7].astype(int)
st = st[st['월'] <= 8]
g1 = st.groupby(['연', '구분', '열차종류', '승하차', '시간대'])['인원'].sum().reset_index()
g2 = st.groupby(['연', '구분', '승하차', '시간대'])['인원'].sum().reset_index(); g2['열차종류'] = '전체'
ga = pd.concat([g1, g2])
ga['키'] = ga['연'].astype(str) + '|' + ga['구분'] + '|' + ga['열차종류'] + '|' + ga['승하차'] + '|' + ga['시간대']
dE = data_sheet('d_역'); write_df(dE, ga[['키', '인원']])
nE = len(ga) + 1

# d_버스 (노선 × 구분 × 시)
bus = pd.read_excel(ROOT / '조사/01_시내버스_노선시간표.xlsx', sheet_name='출발시각_전체')
bus['노선번호'] = bus['노선번호'].astype(str)
bus['시'] = pd.to_numeric(bus['시(時)'], errors='coerce')
bus['시각'] = bus['출발시각'].astype(str).str[:5]
bus['시각값'] = bus['시각'].map(lambda t: int(t[:2]) / 24 + int(t[3:5]) / 1440 if re.match(r'\d\d:\d\d', t) else None)
bg = bus.groupby(['노선번호', '시간표 구분', '시']).size().reset_index(name='회')
bg['키'] = bg['노선번호'] + '|' + bg['시간표 구분'] + '|' + bg['시'].astype(int).astype(str)
dB = data_sheet('d_버스'); write_df(dB, bg[['키', '회']])
nB = len(bg) + 1
dB2 = data_sheet('d_버스출발'); write_df(dB2, bus[['노선번호', '시간표 구분', '시각값']])
nB2 = len(bus) + 1
routes = sorted(bus['노선번호'].unique(), key=lambda x: (not x.isdigit(), int(x) if x.isdigit() else 0, x))
dR = data_sheet('d_노선목록'); write_df(dR, pd.DataFrame({'노선': routes}))
seg = pd.read_excel(ROOT / '조사/01_시내버스_노선시간표.xlsx', sheet_name='시간대별_운행(평일)')
seg = seg.rename(columns={seg.columns[0]: '구간'})
seg = seg[seg['구간'].astype(str).str.contains('→')].reset_index(drop=True)
hours = [f'{h:02d}시' for h in range(5, 24)]
stops = pd.read_excel(ROOT / '조사/04_BIS정류장_위치.xlsx', sheet_name='안동_전체_정류장')
s112 = pd.read_excel(ROOT / '조사/01_시내버스_노선시간표.xlsx', sheet_name='112번_정류장순서')
s112 = s112.merge(stops[['BIS stopId', '위도', '경도']], on='BIS stopId', how='left').dropna(subset=['위도'])
reg = pd.read_excel(ROOT / '조사/04_BIS정류장_위치.xlsx', sheet_name='권역별_정류장')
dP = data_sheet('d_정류장')
write_df(dP, pd.DataFrame({'경도_전체': stops['경도'], '위도_전체': stops['위도'],
                           '경도_안내기': np.where(stops['안내기 설치 여부'] == '설치', stops['경도'], np.nan),
                           '위도_안내기': np.where(stops['안내기 설치 여부'] == '설치', stops['위도'], np.nan)}))
nP = len(stops) + 1
write_df(dP, pd.DataFrame({'경도_112': s112['경도'], '위도_112': s112['위도']}), 1, 6)
n112 = len(s112) + 1
for k, (rg, cc) in enumerate((('안동역', 8), ('원도심', 10), ('월영교', 12))):
    rr = reg[reg['권역'] == rg]
    write_df(dP, pd.DataFrame({f'경도_{rg}': rr['경도'], f'위도_{rg}': rr['위도']}), 1, cc)
nReg = {rg: len(reg[reg['권역'] == rg]) + 1 for rg in ('안동역', '원도심', '월영교')}

# d_월영교 · d_철도
wy = pd.DataFrame(Q['월영교조사'])
wy['업종'] = wy['상권업종중분류명'].str.strip().replace({'비알코올': '카페'})
wy['종료값'] = wy['영업종료'].map(lambda t: int(t[:2]) + int(t[3:5]) / 60 if isinstance(t, str) and re.match(r'\d\d:\d\d', t) else np.nan)
wy = wy[wy['상태'] == '영업'].sort_values('종료값').reset_index(drop=True)
rail = pd.read_csv(ROOT / 'data/external/철도공사_8대도시/안동_관광지_방문소비갭.csv')
dW = data_sheet('d_철도'); write_df(dW, rail[['관광지명', '방문점유율_pct', '방문객소비건수점유율_pct', '소비방문배율']])
nW = len(rail) + 1

# d_진단
dg = pd.read_csv(ROOT / '보고서/전국진단_20260923/전국진단표.csv')
dg = dg.rename(columns={dg.columns[0]: '시군'})
IND = ['①체험문화_변화율', '②평균숙박일수', '③저녁전환율', '④야간방문비중']

# ═════════════ 홈 ═════════════
ws = W['홈']
setup(ws, '안동 이어드림 · 데이터 대시보드', '원도심에서 도는 돈과 월영교에 모이는 사람을 밤까지 잇는 3단계 릴레이: 데이터로 진단하고, 효과를 예측하고, 검증을 설계했다')
label(ws, 'B7', '이야기 흐름 (칸을 누르면 해당 화면으로 이동)', 11, True, NAVY)
tiles = [('1_문제', '① 문제', '방문은 +5.6%인데\n방문당 체험·문화 소비 −16.8%\n(전국 중앙 −2.0%)'),
         ('2_혜택배치', '② 혜택이 어긋남', '방문 1위 중구동\n주민증 혜택업체 0곳\n(도산면 2.4%에 6곳)'),
         ('3_운영시간', '③ 운영 시간 창', '안동역 막차 21~22시\n→ 릴레이 18:30~21:00'),
         ('4_버스', '④ 밤에 끊기는 길', '원도심→월영교\n19시 이후 버스 0회\n(112번 막차 18:45)'),
         ('5_월영교', '⑤ 월영교의 빈자리', '21시까지 식사 3곳/16곳\n주점 0곳\n소비 배율 0.40'),
         ('6_기대효과', '⑥ 기대효과', f"확대안 체험 결제\n{SIM['확대안']['체험결제합']['P50']:,.0f}건\n체험·문화 격차의 {SIM['확대안']['격차기여율']['P50'] * 100:.0f}%"),
         ('7_사후검증', '⑦ 사후 검증', '추첨 순서로 열어 비교\n83곳이면 5개월 안에\n+50% 효과부터 판정'),
         ('8_전국진단', '⑧ 전국 확산', '144개 시·군 자가 진단\n안동과 같은 유형 7곳')]
for i, (name, head, body) in enumerate(tiles):
    r = 9 + (i // 4) * 6; c = 2 + (i % 4) * 4 + (i % 4) // 4
    col = 2 + (i % 4) * 4 + (i % 4)
    ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=col + 3)
    ws.merge_cells(start_row=r + 1, start_column=col, end_row=r + 4, end_column=col + 3)
    h = ws.cell(r, col, head); h.font = F(11, True, 'FFFFFF'); h.fill = FILL(TAB[name]); h.alignment = CENTER; h.hyperlink = f"#{q_(name)}!A1"
    b = ws.cell(r + 1, col, body); b.font = F(10.5, True, INK); b.alignment = CENTER; b.hyperlink = f"#{q_(name)}!A1"
    for rr in range(r + 1, r + 5): ws.row_dimensions[rr].height = 17
    for rr in range(r + 1, r + 5):
        for cc in range(col, col + 4):
            ws.cell(rr, cc).fill = FILL(LIGHT)
    for cc in range(col, col + 4):
        ws.cell(r, cc).fill = FILL(TAB[name])
label(ws, 'B22', '쓰는 법', 11, True, NAVY)
for i, t in enumerate(['● 노란 칸 = 직접 바꿔 보는 값. 목록(▼)을 고르거나 숫자를 넣으면 그 화면의 숫자 카드와 차트가 바로 다시 계산됩니다.',
                       '● 위쪽 파란 메뉴를 누르면 다른 화면으로 이동합니다. 원자료와 계산 시트는 숨겨 두었습니다(시트 탭 오른쪽 클릭 → 숨기기 취소로 볼 수 있음).',
                       '● 기대효과 화면은 1만 번이 아니라 3,000번 모의실험을 엑셀 안에서 다시 돌립니다(F9 = 다시 뽑기). 보고서 숫자는 파이썬 1만 회 결과입니다.',
                       '● 모든 예상 효과는 시행 전 조건부 값입니다. 실제 효과는 「사후 검증」 화면의 방법으로 시행 후에 판정합니다.']):
    note(ws, 23 + i, t, 20, size=10, color=INK)
label(ws, 'B28', '결과물 7개와 화면', 11, True, NAVY)
deliv = [('결과물', '화면', '상태'),
         ('① 식음→체험 연결(혜택 배치 진단)', '혜택 배치', '완료 · 체험 연결 경로표는 현장조사 후'),
         ('② 운영 시간 창', '운영 시간', '완료'),
         ('③ 교통 공백 확인', '버스(BIS)', '버스 완료 · 택시 실측 9/26'),
         ('④ 야간 팝업(영업 종료·업종 공백)', '월영교 야간', '완료 · 현장조사 조사일 기록 보완'),
         ('⑤ 행동검증·예상 효과', '기대효과', '완료 · 9/26 관측값으로 갱신'),
         ('⑥ 사후 효과 검증계획', '사후 검증', '완료'),
         ('⑦ 전국 진단표', '전국 진단', '완료'),
         (f"(보조) 팀 온라인 설문 {SVY['표본']['응답']}명(방문 경험 {SVY['표본']['3년내_방문_예']}명)", '방문객 설문', '완료 · 9/23~24 구글폼, 의향은 참여율 아님')]
links = [None, '2_혜택배치', '3_운영시간', '4_버스', '5_월영교', '6_기대효과', '7_사후검증', '8_전국진단', '10_설문']
for k, (a, b, c) in enumerate(deliv):
    r = 29 + k
    for (c1, c2, v) in ((2, 8, a), (9, 11, b), (12, 19, c)):
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        cell = ws.cell(r, c1, v); cell.font = F(9.5, k == 0, NAVY if (k and c1 == 9) else INK); cell.alignment = LEFT if c1 != 9 else CENTER
        for cc in range(c1, c2 + 1):
            ws.cell(r, cc).border = BOX
            if k == 0: ws.cell(r, cc).fill = FILL('D9E1F2')
        if k and c1 == 9: cell.hyperlink = f"#{q_(links[k])}!A1"

# ═════════════ 1 문제 ═════════════
ws = W['1_문제']
setup(ws, '문제 · 방문은 늘었는데 체험·문화 소비는 줄었다', '관광객 감소도, 당일치기 증가도 아니다. 방문당 체험·문화 소비만 전국보다 크게 줄었다(2024→2026년 1~8월, 외지인)')
ref_c, cell_c = control(ws, 7, 2, '비교할 시·군 ▼', '경주시', options=f"='d_시군'!$A$2:$A${nS}")
note(ws, 7, '← 목록에서 시·군을 고르면 파란 카드와 차트의 파란 막대가 그 지역으로 바뀝니다(빨강 = 안동)', col1=6, col2=17, size=10, color=BLUE)
colmap = {c: CL(i + 1) for i, c in enumerate(sg.columns)}
look = lambda col: f"INDEX('d_시군'!${col}$2:${col}${nS},MATCH({ref_c},'d_시군'!$A$2:$A${nS},0))"
andong = lambda col: f"INDEX('d_시군'!${col}$2:${col}${nS},MATCH(\"안동시\",'d_시군'!$A$2:$A${nS},0))"
MEDF = lambda col: f"MEDIAN('d_시군'!${col}$2:${col}${nS})"
PCT = '+0.0%;-0.0%'
IND1 = [('외지인 방문 변화율', colmap['방문_변화율']), ('방문당 관광소비 변화율', colmap['관광소비_변화율']), ('방문당 체험·문화 소비 변화율', colmap['체험문화_변화율'])]
for j, (lab, col) in enumerate(IND1):
    kpi(ws, 9, 2 + 4 * j, f'안동 · {lab}', f"={andong(col)}", PCT, f'="전국 중앙 "&TEXT({MEDF(col)},"{PCT}")', color=RED)
kpi(ws, 9, 14, '안동 · 체험·문화 감소폭 순위', f"=MATCH(\"안동시\",'d_시군'!$A$2:$A${nS},0)", '0"위"', f'{len(sg)}개 시·군 중(1위 = 가장 많이 줄어듦)', color=RED)
for j, (lab, col) in enumerate(IND1):
    kpi(ws, 14, 2 + 4 * j, f'="선택 · "&{ref_c}&" · {lab}"', f"={look(col)}", PCT,
        f'="안동과 차이 "&TEXT(({look(col)}-{andong(col)})*100,"+0.0;-0.0")&"%p"', color=BLUE)
kpi(ws, 14, 14, f'="선택 · "&{ref_c}&" · 감소폭 순위"', f"=MATCH({ref_c},'d_시군'!$A$2:$A${nS},0)", '0"위"', f'{len(sg)}개 시·군 중', color=BLUE)
# 차트 1: 146곳 체험·문화 변화율(정렬), 안동 빨강·선택 파랑
cE = colmap['체험문화_변화율']
dS.cell(1, 12, '전체'); dS.cell(1, 13, '안동'); dS.cell(1, 14, '선택')
for i in range(2, nS + 1):
    dS.cell(i, 12, f"={cE}{i}")
    dS.cell(i, 13, f'=IF(A{i}="안동시",{cE}{i},"")')
    dS.cell(i, 14, f'=IF(A{i}={ref_c},{cE}{i},"")')
ch = BarChart(); ch.type = 'col'; ch.grouping = 'clustered'; ch.overlap = 100; ch.gapWidth = 20
for col, colr, t in ((12, 'D0D0D0', '다른 시·군'), (13, RED, '안동'), (14, BLUE, '선택 시·군')):
    sr = Series(Reference(dS, min_col=col, min_row=2, max_row=nS), title=t); color_series(sr, colr); ch.series.append(sr)
ch.set_categories(Reference(dS, min_col=1, min_row=2, max_row=nS)); ch.x_axis.delete = True; ch.y_axis.number_format = '0%'
ch.y_axis.scaling.min = -0.6; ch.y_axis.scaling.max = 0.8; ch.y_axis.majorUnit = 0.2
style_chart(ch, f'{len(sg)}개 시·군의 방문당 체험·문화 소비 변화율(2024→2026년 1~8월, 왼쪽일수록 많이 줄어듦)', 34, 8)
ws.add_chart(ch, 'B19')
# 업종: 안동 vs 선택 (방문당 소비, 전국 중앙 추세 대비 차이)
nU = min(len(up), 12)
dU.append(['업종', '안동_2024', '안동_2026', '전국중앙_변화율', '선택_2024', '선택_2026', '안동 중앙대비(원)', '선택 중앙대비(원)'])
for i in range(2, nU + 2):
    r = up.iloc[i - 2]
    dU.cell(i, 1, r['업종']); dU.cell(i, 2, float(r['2024'])); dU.cell(i, 3, float(r['2026'])); dU.cell(i, 4, float(r['전국중앙_변화율']))
    for c_, yc in ((5, 'B'), (6, 'C')):
        dU.cell(i, c_, f"=SUMIFS('d_업종전체'!${yc}$2:${yc}${nUa},'d_업종전체'!$A$2:$A${nUa},{ref_c}&\"|\"&A{i})")
    dU.cell(i, 7, f"=C{i}-B{i}*(1+D{i})"); dU.cell(i, 8, f"=F{i}-E{i}*(1+D{i})")
dU.cell(1, 10, '선택 이름'); dU.cell(2, 10, f"={ref_c}")
ch = BarChart(); ch.type = 'bar'; ch.grouping = 'clustered'; ch.gapWidth = 40
for col, colr, t in ((3, RED, '안동'), (6, BLUE, '선택 시·군')):
    sr = Series(Reference(dU, min_col=col, min_row=2, max_row=nU + 1), title=t); color_series(sr, colr); ch.series.append(sr)
ch.set_categories(Reference(dU, min_col=1, min_row=2, max_row=nU + 1)); ch.x_axis.scaling.orientation = 'maxMin'; ch.y_axis.number_format = '#,##0'
style_chart(ch, '업종별 외지인 방문당 소비(원, 2026년 1~8월): 체험·문화(문화서비스·기타레저)는 아래쪽, 소비의 약 1%', 17, 10.5)
ws.add_chart(ch, 'B36')
ch = BarChart(); ch.type = 'bar'; ch.grouping = 'clustered'; ch.gapWidth = 40
for col, colr, t in ((7, RED, '안동'), (8, BLUE, '선택 시·군')):
    sr = Series(Reference(dU, min_col=col, min_row=2, max_row=nU + 1), title=t); color_series(sr, colr); ch.series.append(sr)
ch.set_categories(Reference(dU, min_col=1, min_row=2, max_row=nU + 1)); ch.x_axis.scaling.orientation = 'maxMin'; ch.x_axis.tickLblPos = 'low'
ch.y_axis.number_format = '#,##0'
style_chart(ch, '전국 중앙 추세 대비 방문당 소비 차이(원): −는 전국만큼 늘지 못한 업종', 17, 10.5)
ws.add_chart(ch, 'K36')
note(ws, 57, '읽는 법: 관광객 감소(방문 +5.6%)도, 당일치기 증가(무박 86.3→86.4%, 2019→2025)도 아니다. 방문당 체험·문화 소비는 −16.8%로 전국 146곳 중 감소폭 22번째다. '
     '다만 체험·문화는 관광소비의 약 1%이고, 하락의 3분의 2는 문화서비스가 2025년 중반 한 번에 꺾인 것이라 원인을 한 가지로 단정하지 않는다. 외식·숙박은 전국보다 잘 되고 있다. '
     '오른쪽 차트의 "전국 중앙 추세 대비 차이" = 2026년 값 − 2024년 값 × (1 + 그 업종의 146곳 중앙 변화율).', 48)
note(ws, 58, '출처: 한국관광 데이터랩 외지인 신용카드 소비(BDT_02_01_003, 업종 중분류) ÷ 외지인 방문 연인원(BDT_01_01_006), 2024·2026년 1~8월. 체험·문화 = 문화서비스 + 관광유원시설 + 기타레저.', 30, size=8.5, color='7F7F7F')

# ═════════════ 2 혜택 배치 ═════════════
ws = W['2_혜택배치']
setup(ws, '혜택 배치 · 사람이 모이는 곳에 혜택이 없다', '외지인 방문 1위 중구동(원도심)에 주민증 혜택업체 0곳, 12위 도산면에 6곳. 혜택을 옮기면 어긋남이 얼마나 줄어드는지 바꿔 보세요')
ref_d, _ = control(ws, 7, 2, '혜택을 추가할 동 ▼', '중구동', options=f"='d_혜택'!$A$2:$A${len(hb) + 1}")
ref_n, _ = control(ws, 7, 7, '추가할 업체 수', 0, '0"곳"')
note(ws, 7, '← 동을 고르고 업체 수를 넣으면 카드와 차트(빨강)에 반영됩니다', col1=10, col2=17, size=9)
nh = len(hb) + 1
dH.cell(1, 4, '추가'); dH.cell(1, 5, '혜택 합계'); dH.cell(1, 6, '혜택 점유율'); dH.cell(1, 7, '방문 점유율(비율)')
for i in range(2, nh + 1):
    dH.cell(i, 4, f"=IF(A{i}={ref_d},{ref_n},0)")
    dH.cell(i, 5, f"=C{i}+D{i}")
    dH.cell(i, 6, f"=E{i}/SUM($E$2:$E${nh})")
    dH.cell(i, 7, f"=B{i}/100")
kpi(ws, 9, 2, '중구동 혜택업체(추가 반영)', f"=INDEX('d_혜택'!$E$2:$E${nh},MATCH(\"중구동\",'d_혜택'!$A$2:$A${nh},0))", '0"곳"', '추가 전 0곳 (방문 점유율 12.5%, 1위)', color=RED)
kpi(ws, 9, 6, '방문 상위 5개 동의 혜택 비중', f"=SUM('d_혜택'!$E$2:$E$6)/SUM('d_혜택'!$E$2:$E${nh})", '0%', f'="추가 전 "&TEXT(SUM(\'d_혜택\'!$C$2:$C$6)/SUM(\'d_혜택\'!$C$2:$C${nh}),"0%")&" · 방문 비중은 "&TEXT(SUM(\'d_혜택\'!$G$2:$G$6),"0%")')
kpi(ws, 9, 10, '방문과 혜택의 어긋남 지수', f"=0.5*SUMPRODUCT(ABS('d_혜택'!$G$2:$G${nh}-'d_혜택'!$F$2:$F${nh}))", '0%', f'="추가 전 "&TEXT(0.5*SUMPRODUCT(ABS(\'d_혜택\'!$G$2:$G${nh}-\'d_혜택\'!$C$2:$C${nh}/SUM(\'d_혜택\'!$C$2:$C${nh}))),"0%")&" · 0% = 방문 비중대로 배치"')
kpi(ws, 9, 14, '주민증 이용 중 관람지 비중', P['A_업체'] and sum(r['이용'] for r in P['A_업체'] if r['분류'] == '관람') / sum(r['이용'] for r in P['A_업체']), '0.0%',
    '체험 0.9% (월 약 7건) · 이용 표시 합계 20,477건')
ch = BarChart(); ch.type = 'col'; ch.grouping = 'clustered'
for col, t, colr in ((7, '외지인 방문 점유율', GRAY), (6, '혜택업체 점유율(추가 반영)', RED)):
    s = Series(Reference(dH, min_col=col, min_row=2, max_row=nh), title=t); color_series(s, colr); ch.series.append(s)
ch.set_categories(Reference(dH, min_col=1, min_row=2, max_row=nh)); ch.y_axis.number_format = '0%'
style_chart(ch, '읍면동별 외지인 방문 점유율 vs 주민증 혜택업체 점유율 (2026년 1~8월, 방문 순)', 34, 9)
ws.add_chart(ch, 'B14')
cls = pd.DataFrame(P['A_업체']).groupby('분류')['이용'].sum().sort_values(ascending=False).reset_index()
r0 = 33
label(ws, f'B{r0}', '주민증 가맹 업체별 이용 표시(누적, 2024.6~2026.9)', 10.5, True, NAVY)
tb = bz[['혜택업체명', '분류', '행정동', '이용', '이용률_1만']].rename(columns={'혜택업체명': '업체', '이용': '이용 표시(건)', '이용률_1만': '방문 1만 회당 월 이용'})
tcols = [(2, 3), (4, 4), (5, 5), (6, 6), (7, 8)]                  # 업체명·이용률은 두 칸 병합
for i, row in enumerate([list(tb.columns)] + tb.values.tolist()):
    r = r0 + 1 + i
    for (c1, c2), v, fm in zip(tcols, row, (None, None, None, '#,##0', '0.00')):
        if c2 > c1: ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        cell = ws.cell(r, c1, v); cell.font = F(9, i == 0); cell.alignment = CENTER if (i == 0 or c1 > 2) else LEFT
        if i == 0: cell.fill = FILL('D9E1F2')
        if fm and i: cell.number_format = fm
        for cc in range(c1, c2 + 1): ws.cell(r, cc).border = BOX
dC = data_sheet('d_분류'); write_df(dC, cls)
ch = BarChart(); ch.type = 'bar'
s = Series(Reference(dC, min_col=2, min_row=2, max_row=len(cls) + 1), title='이용 표시(건)'); color_series(s, NAVY); ch.series.append(s)
ch.set_categories(Reference(dC, min_col=1, min_row=2, max_row=len(cls) + 1)); ch.x_axis.scaling.orientation = 'maxMin'
style_chart(ch, '분류별 주민증 이용: 관람지 95%, 체험 0.9%', 15, 8, None)
ws.add_chart(ch, 'I34')
note(ws, r0 + len(tb) + 3, '읽는 법: 혜택은 유교문화 관람지에 몰려 있고, 사람이 가장 많은 원도심에는 없다. 업체 이용도 관람지(하회마을·도산서원·유교랜드)에 95%가 몰린다. '
     '주민증 가맹 26곳 회귀 결과, 업체가 있는 동에 방문이 많아도 이용은 거의 늘지 않았다(탄력성 0.31, 유의하지 않음) → 혜택을 옮기는 것만으로는 부족하고 체험으로 잇는 장치가 필요하다.', 44)
sv11 = SVY['Q11_유료체험_망설임']; n11 = SVY['Q11_비가격_이유']['n']
note(ws, r0 + len(tb) + 5, f"방문객 설문({n11}명 응답): 유료 체험을 망설이는 이유는 정보 부족 {sv11['체험 정보가 부족해서'] / n11:.0%} · 줄·예약 불편 {sv11['줄을 오래 서거나 예약하기 불편해서'] / n11:.0%} · "
     f"이동 시간 부족 {sv11['다음 장소로 이동할 시간이 부족해서'] / n11:.0%} · 가격 {sv11['가격이 부담스러워서'] / n11:.0%} → 혜택은 할인만이 아니라 체험 정보·예약을 잇는 쪽으로. 자세히는 「방문객 설문」 화면.", 30, color=NAVY)
note(ws, r0 + len(tb) + 4, '출처: 데이터랩 읍면동별 외지인 방문자 수(BDT_01_01_005_1) · 안동시 정보공개청구 회신(혜택업체 27곳, 2026-09-18) · 한국관광공사 주민증 누리집 가맹점 이용 표시(2026-09-21 조회) · 소상공인 상가정보(행정동 대조)', 30, size=8.5, color='7F7F7F')

# ═════════════ 3 운영 시간 ═════════════
ws = W['3_운영시간']
setup(ws, '운영 시간 창 · 안동역 막차가 릴레이의 끝을 정한다', '주말 저녁 승차는 18~19시에 가장 많고, 막차는 21~22시다. 릴레이 시작·종료 시각을 바꿔 창 안에 들어오는 승차를 확인해 보세요')
ref_y, _ = control(ws, 7, 2, '연도 ▼', '2026', options='"2024,2025,2026"')
ref_g, _ = control(ws, 7, 5, '요일 구분 ▼', '주말', options='"평일,주말,공휴일,명절대수송"')
ref_t, _ = control(ws, 7, 8, '열차 ▼', '전체', options='"전체,KTX-이음,새마을,무궁화"')
ref_s, _ = control(ws, 7, 11, '릴레이 시작(시)', 18.5, '0.0')
ref_e, _ = control(ws, 7, 14, '릴레이 종료(시)', 21, '0.0')
dT = data_sheet('d_역계산')
dT.append(['시간대', '시', '승차', '하차', '창 안 승차'])
for i, h in enumerate(range(5, 24), 2):
    key = f'{h:02d}-{h + 1:02d}'
    dT.cell(i, 1, f'{h:02d}시'); dT.cell(i, 2, h)
    for col, sb in ((3, '승차'), (4, '하차')):
        dT.cell(i, col, f"=SUMIFS('d_역'!$B$2:$B${nE},'d_역'!$A$2:$A${nE},{ref_y}&\"|\"&{ref_g}&\"|\"&{ref_t}&\"|{sb}|{key}\")")
    dT.cell(i, 5, f'=IF(AND(B{i}+1>{ref_s},B{i}<{ref_e}),C{i},"")')
nT = 20
kpi(ws, 9, 2, '저녁(17~22시) 최대 승차 시간대', f"=INDEX('d_역계산'!$A$14:$A$18,MATCH(MAX('d_역계산'!$C$14:$C$18),'d_역계산'!$C$14:$C$18,0))", '@',
    f"=\"승차 \"&TEXT(MAX('d_역계산'!$C$14:$C$18),\"#,##0\")&\"명 (1~8월 합계)\"")
kpi(ws, 9, 6, '마지막 승차 시간대(막차)', "=LOOKUP(2,1/('d_역계산'!$C$2:$C$20>0),'d_역계산'!$A$2:$A$20)", '@', "=\"22시 이후 승차 \"&TEXT(SUM('d_역계산'!$C$19:$C$20),\"#,##0\")&\"명\"", color=RED)
kpi(ws, 9, 10, '릴레이 창 안 승차 비중', "=SUMIFS('d_역계산'!$C$2:$C$20,'d_역계산'!$B$2:$B$20,\">=\"&INT(" + ref_s + "),'d_역계산'!$B$2:$B$20,\"<\"&" + ref_e + ")/SUM('d_역계산'!$C$2:$C$20)", '0.0%',
    f'="창 "&TEXT({ref_s}/24,"hh:mm")&"~"&TEXT({ref_e}/24,"hh:mm")')
kpi(ws, 9, 14, '종료 후 막차까지 여유', f"=MAX(0,LOOKUP(2,1/('d_역계산'!$C$2:$C$20>0),'d_역계산'!$B$2:$B$20)+1-{ref_e})", '0.0"시간"',
    '=IF(' + ref_e + '>LOOKUP(2,1/(\'d_역계산\'!$C$2:$C$20>0),\'d_역계산\'!$B$2:$B$20)+1,"막차 이후 → 귀가 수단 없음",IF(' + ref_e + '>LOOKUP(2,1/(\'d_역계산\'!$C$2:$C$20>0),\'d_역계산\'!$B$2:$B$20),"막차 시간대와 겹침 → 여유 없음","막차 전에 끝남"))')
ch = BarChart(); ch.type = 'col'; ch.grouping = 'clustered'; ch.overlap = 100; ch.gapWidth = 40
for col, t, colr in ((3, '승차', 'C9C9C9'), (5, '릴레이 창 안 승차', RED)):
    s = Series(Reference(dT, min_col=col, min_row=2, max_row=nT), title=t); color_series(s, colr); ch.series.append(s)
ch.set_categories(Reference(dT, min_col=1, min_row=2, max_row=nT)); ch.y_axis.number_format = '#,##0'
style_chart(ch, '안동역 시간대별 승차(선택한 연도·요일·열차, 1~8월 합계)', 34, 9)
ws.add_chart(ch, 'B14')
ch = LineChart()
s = Series(Reference(dT, min_col=4, min_row=2, max_row=nT), title='하차'); color_series(s, NAVY, True); ch.series.append(s)
ch.set_categories(Reference(dT, min_col=1, min_row=2, max_row=nT)); ch.y_axis.number_format = '#,##0'
style_chart(ch, '안동역 시간대별 하차(도착)', 34, 7, None)
ws.add_chart(ch, 'B33')
note(ws, 48, '읽는 법: 시간 분포는 수요가 아니라 열차 시각표를 반영한다(22시 이후 0 = 열차가 없음). 그래서 막차가 설계의 제약 조건이 되고, 릴레이는 18:30에 시작해 21:00에 끝나야 막차 전에 닿는다. '
     '승차 인원에는 주민이 섞여 있어 "관광객"이라 부르지 않는다. 값은 그 달 해당 요일 전체의 합이다.', 40)
q13 = SVY['Q13_당일귀가자중_20~22시']
note(ws, 50, f"방문객 설문: 당일 귀가한 방문자 {q13['n']}명 중 {q13['k']}명({q13['pct']:.0f}%)이 20~22시 막차 시간대에 귀가했다고 답했다 → 릴레이가 21시 전에 끝나야 하는 이유와 같은 방향. 자세히는 「방문객 설문」 화면.", 30, color=NAVY)
note(ws, 49, '출처: 코레일 안동역 월별·요일별·시간대별 승하차(KTX-이음·새마을·무궁화, 2024.1~2026.8, 팀원 취합), 연도 간 비교를 위해 1~8월만 합산', 20, size=8.5, color='7F7F7F')

# ═════════════ 4 버스 ═════════════
ws = W['4_버스']
setup(ws, '버스(BIS) · 원도심에서 월영교로 가는 밤길이 끊긴다', '안동역↔원도심은 밤 10시까지 다니지만 원도심→월영교는 112번 한 노선, 하루 6회, 19시 이후 0회다. 노선을 고르고 저녁 운행을 추가해 보세요')
ref_r, _ = control(ws, 7, 2, '노선 ▼', '112', options=f"='d_노선목록'!$A$2:$A${len(routes) + 1}")
ref_k, _ = control(ws, 7, 5, '시간표 ▼', '평일', options='"평일,공휴일(일요일)"')
ref_x, _ = control(ws, 7, 8, '19~21시 원도심→월영교 추가 운행', 0, '0"회"', width=2)
note(ws, 7, '← 추가 운행 횟수를 넣으면 카드와 빨간 선에 반영됩니다', col1=12, col2=18, size=9)
dBc = data_sheet('d_버스계산')
dBc.append(['시', '선택 노선'] + list(seg['구간']) + ['원도심→월영교(추가 반영)'])
for i, h in enumerate(range(5, 24), 2):
    dBc.cell(i, 1, f'{h:02d}시')
    dBc.cell(i, 2, f"=SUMIFS('d_버스'!$B$2:$B${nB},'d_버스'!$A$2:$A${nB},{ref_r}&\"|\"&{ref_k}&\"|{h}\")")
    for j, sname in enumerate(seg['구간']):
        v = seg.loc[seg['구간'] == sname, f'{h:02d}시'].values
        dBc.cell(i, 3 + j, float(v[0]) if len(v) and pd.notna(v[0]) else 0)
    jw = list(seg['구간']).index('원도심→월영교')
    dBc.cell(i, 3 + len(seg), f"={CL(3 + jw)}{i}+IF(AND({h}>=19,{h}<21),{ref_x}/2,0)")
dB2.cell(1, 4, '선택 노선 시각')                                   # 첫차·막차용 보조 열(MINIFS 없는 엑셀 2016에서도 동작)
for i in range(2, nB2 + 1):
    dB2.cell(i, 4, f'=IF(AND(A{i}&""={ref_r}&"",B{i}={ref_k}),C{i},"")')
kpi(ws, 9, 2, '선택 노선 운행 횟수', "=SUM('d_버스계산'!$B$2:$B$20)", '0"회"',
    f"=IF(SUM('d_버스계산'!$B$2:$B$20)=0,\"이 시간표에는 운행 없음\",\"첫차 \"&TEXT(MIN('d_버스출발'!$D$2:$D${nB2}),\"hh:mm\")&\" · 막차 \"&TEXT(MAX('d_버스출발'!$D$2:$D${nB2}),\"hh:mm\"))")
kpi(ws, 9, 6, '원도심→월영교 19시 이후 운행(추가 반영)', f"=SUM('d_버스계산'!${CL(3 + len(seg))}$16:${CL(3 + len(seg))}$20)", '0"회"', '추가 전 0회 · 112번 원도심 출발 막차 18:45(기점)', color=RED)
kpi(ws, 9, 10, '안동역→원도심 19시 이후 운행', f"=SUM('d_버스계산'!$C$16:$C$20)", '0"회"', '막차 22:10 · 원도심→안동역 막차 22:40')
kpi(ws, 9, 14, '버스정보 안내기 설치 정류장', int((stops['안내기 설치 여부'] == '설치').sum()), '#,##0"곳"', f'전체 {len(stops):,}곳 중 · 월영교 500m 4곳 중 1곳')
ch = BarChart(); ch.type = 'col'
s = Series(Reference(dBc, min_col=2, min_row=2, max_row=20), title='선택 노선 출발'); color_series(s, NAVY); ch.series.append(s)
ch.set_categories(Reference(dBc, min_col=1, min_row=2, max_row=20))
style_chart(ch, '선택 노선의 시간대별 기점 출발 횟수', 17, 8.5, None)
ws.add_chart(ch, 'B14')
ch = LineChart()
cols_seg = [(3 + list(seg['구간']).index(n), n, c) for n, c in (('안동역→원도심', GRAY), ('원도심→안동역', 'D0D0D0'))]
for col, t, colr in cols_seg:
    s = Series(Reference(dBc, min_col=col, min_row=2, max_row=20), title=t); color_series(s, colr, True); ch.series.append(s)
s = Series(Reference(dBc, min_col=3 + len(seg), min_row=2, max_row=20), title='원도심→월영교(추가 반영)'); color_series(s, RED, True); ch.series.append(s)
ch.set_categories(Reference(dBc, min_col=1, min_row=2, max_row=20))
style_chart(ch, '구간별 시간대 운행 횟수(평일, 기점 출발)', 17, 8.5)
ws.add_chart(ch, 'K14')
sc = ScatterChart(); sc.style = 13
def add_pts(xc, yc, n, title, colr, size):
    s = Series(Reference(dP, min_col=yc, min_row=2, max_row=n), Reference(dP, min_col=xc, min_row=2, max_row=n), title=title)
    s.marker.symbol = 'circle'; s.marker.size = size; s.marker.graphicalProperties.solidFill = colr; s.marker.graphicalProperties.line.solidFill = colr
    s.graphicalProperties.line.noFill = True; sc.series.append(s)
add_pts(1, 2, nP, '정류장 전체', 'D9D9D9', 3)
add_pts(3, 4, nP, '안내기 설치', BLUE, 4)
add_pts(6, 7, n112, '112번 노선(월영교 유일)', RED, 5)
for rg, cc, colr in (('안동역', 8, '000000'), ('원도심', 10, 'ED7D31'), ('월영교', 12, '7030A0')):
    add_pts(cc, cc + 1, nReg[rg], f'{rg} 권역', colr, 7)
sc.x_axis.scaling.min = 128.64; sc.x_axis.scaling.max = 128.80; sc.y_axis.scaling.min = 36.53; sc.y_axis.scaling.max = 36.62
sc.x_axis.title = '경도'; sc.y_axis.title = '위도'; sc.x_axis.number_format = '0.00'; sc.y_axis.number_format = '0.00'
style_chart(sc, '안동 시내 버스정류장 지도(좌표): 112번이 원도심과 월영교를 잇는 유일한 노선', 34, 13)
ws.add_chart(sc, 'B32')
note(ws, 59, '읽는 법: 버스 시각은 기점 출발 시각이라 구간 통과는 더 늦다. 공휴일 시간표가 없는 노선이 많아 토요일 운행은 단정하지 않는다. 추가 운행 칸은 "관광택시 저녁 재배치나 셔틀로 몇 회를 채우면" 그래프가 어떻게 바뀌는지 보는 가정이다.', 34)
q3 = SVY['Q3_버스불편_버스이용방문자']; q4w = SVY['Q4_포기한곳_방문자']['야간 명소(월영교)']
note(ws, 61, f"방문객 설문: 시내버스로 다닌 방문자 {q3['n']}명 중 {q3['k']}명({q3['pct']:.0f}%)이 '불편했다'. 이동이 불편해 포기한 곳 1위는 월영교 같은 야간 명소({q4w['pct']:.0f}%, {q4w['k']}/{q4w['n']}).", 30, color=NAVY)
note(ws, 60, '출처: 안동시 버스정보시스템 누리집 공개 시간표·노선·정류장(2026-09-22 조회, 46개 노선·3,309개 정류장, 안내기 = bitYn), 국토교통부 전국 버스정류장 위치정보(대조)', 20, size=8.5, color='7F7F7F')

# ═════════════ 5 월영교 ═════════════
ws = W['5_월영교']
setup(ws, '월영교 야간 · 사람은 모이는데 쓸 곳이 없다', '외곽 방문자의 54%가 월영교에 들르지만 소비 배율은 0.40, 21시까지 식사 가능한 곳은 16곳 중 3곳, 주점 0곳. 팝업을 넣어 빈자리를 채워 보세요')
ref_b, _ = control(ws, 7, 2, '팝업 식사·주점 부스', 8, '0"개"')
ref_z, _ = control(ws, 7, 5, '팝업 종료 시각(시)', 22, '0')
ref_rl, _ = control(ws, 7, 8, '관광지 ▼', '월영교', options=f"='d_철도'!$A$2:$A${nW}")
dM = data_sheet('d_월영교')
dM.append(['업체', '업종', '영업 종료(시)', '식사 가능'])
for _, r in wy.iterrows():
    dM.append([r['상호명'], r['업종'], r['종료값'], 1 if r['업종'] in ('한식', '서양식') else 0])
nM = len(wy) + 1
dM.cell(nM + 1, 1, '팝업 부스'); dM.cell(nM + 1, 2, '팝업'); dM.cell(nM + 1, 3, f"={ref_z}"); dM.cell(nM + 1, 4, 1)
kpi(ws, 9, 2, '21시까지 식사 가능한 곳(팝업 반영)', f"=COUNTIFS('d_월영교'!$D$2:$D${nM},1,'d_월영교'!$C$2:$C${nM},\">=21\")+IF({ref_z}>=21,{ref_b},0)", '0"곳"',
    f"=\"팝업 전 3곳 · 식사 가능 비중 \"&TEXT((COUNTIFS('d_월영교'!$D$2:$D${nM},1,'d_월영교'!$C$2:$C${nM},\">=21\")+IF({ref_z}>=21,{ref_b},0))/({nM - 1}+{ref_b}),\"0%\")", color=RED)
kpi(ws, 9, 6, '주점(팝업 반영)', f"=IF({ref_z}>=21,{ref_b},0)", '0"곳"', '팝업 전 0곳 (원도심 1km는 764곳 중 86곳)', color=RED)
kpi(ws, 9, 10, '선택 관광지 소비 전환 배율', f"=INDEX('d_철도'!$D$2:$D${nW},MATCH({ref_rl},'d_철도'!$A$2:$A${nW},0))", '0.00',
    f"=\"방문 \"&TEXT(INDEX('d_철도'!$B$2:$B${nW},MATCH({ref_rl},'d_철도'!$A$2:$A${nW},0)),\"0.0\")&\"% vs 소비건수 \"&TEXT(INDEX('d_철도'!$C$2:$C${nW},MATCH({ref_rl},'d_철도'!$A$2:$A${nW},0)),\"0.0\")&\"% (2022.4~6)\"")
kpi(ws, 9, 14, '외곽 방문자 중 월영교도 방문', 0.548, '0.0%', '만휴정 54.8% · 도산서원 54.2% (철도공사 연관규칙)')
ch = BarChart(); ch.type = 'bar'
s = Series(Reference(dM, min_col=3, min_row=2, max_row=nM + 1), title='영업 종료 시각'); color_series(s, GRAY); ch.series.append(s)
ch.set_categories(Reference(dM, min_col=1, min_row=2, max_row=nM + 1)); ch.x_axis.scaling.orientation = 'maxMin'
ch.y_axis.scaling.min = 17; ch.y_axis.scaling.max = 24; ch.y_axis.majorUnit = 1
style_chart(ch, '월영교 반경 1km 영업 음식점의 종료 시각(시) + 팝업(맨 아래)', 17, 10, None)
ws.add_chart(ch, 'B14')
uw = pd.DataFrame({'업종': ['한식', '카페(비알코올)', '주점', '서양식', '기타'],
                   '월영교 1km': [10 / 17, 6 / 17, 0, 1 / 17, 0],
                   '원도심 1km': [342 / 764, 127 / 764, 86 / 764, 13 / 764, (764 - 342 - 127 - 86 - 13) / 764]})
dK = data_sheet('d_업종구성'); write_df(dK, uw)
ch = BarChart(); ch.type = 'col'
for col, colr in ((2, RED), (3, GRAY)):
    s = Series(Reference(dK, min_col=col, min_row=2, max_row=6), title=dK.cell(1, col).value); color_series(s, colr); ch.series.append(s)
ch.set_categories(Reference(dK, min_col=1, min_row=2, max_row=6)); ch.y_axis.number_format = '0%'
style_chart(ch, '음식점 업종 구성: 월영교 1km(17곳) vs 원도심 1km(764곳)', 17, 10)
ws.add_chart(ch, 'K14')
ch = BarChart(); ch.type = 'bar'
for col, colr, t in ((2, GRAY, '방문 점유율(%)'), (3, RED, '소비건수 점유율(%)')):
    s = Series(Reference(dW, min_col=col, min_row=2, max_row=nW), title=t); color_series(s, colr); ch.series.append(s)
ch.set_categories(Reference(dW, min_col=1, min_row=2, max_row=nW)); ch.x_axis.scaling.orientation = 'maxMin'
style_chart(ch, '관광지별 방문 vs 소비 점유율(철도공사 가명결합, 2022년 4~6월): 월영교·문화관광단지·하회마을은 방문만큼 쓰지 않는다', 34, 11)
ws.add_chart(ch, 'B35')
note(ws, 58, '읽는 법: 월영교는 외곽을 돈 방문객이 돌아오며 모이는 곳인데, 반경 1km 식당 대부분이 20:30 전에 닫고 22시 이후엔 카페만 남으며 주점은 없다. 18시 이후 운영하는 체험도 문보트·황포돛배뿐이다. '
     '팝업은 21시 이후 식사·주류를 채우는 자리이며, 부스 수는 처리량·부지 조건 확인 후 정한다(예시 8개 = 2026 월영장터 장터 규모). 참고로 연 10일인 월영야행 기간에도 권역 월 소비가 통계적으로 늘지 않았다 → 상시 운영·원도심 연결이 필요한 이유.', 50)
q9, q9c = SVY['Q9_월영교밤_부족_방문자'], SVY['Q9_월영교밤_충분_방문자']
note(ws, 60, f"방문객 설문: 20시 이후 월영교 주변 식당·편의시설이 '부족·매우 부족' {q9['pct']:.0f}%({q9['k']}/{q9['n']}), '충분' {q9c['pct']:.0f}%({q9c['k']}명). 21~23시 팝업 포차는 방문자 {SVY['Q10_야간팝업포차']['방문자_top2']['pct']:.0f}%가 긍정(의향이지 이용률이 아님).", 30, color=NAVY)
note(ws, 59, '출처: 소상공인 상가(상권)정보(2026-06-30) · 팀 현장조사(영업 종료, 2026-09 확인, 조사일·방법 보완 예정) · 한국철도공사 8대 도시(안동) 관광 형태 분석(2022.4~6, 소비는 건수)', 20, size=8.5, color='7F7F7F')

# ═════════════ 6 기대효과 (모의실험 3,000회) ═════════════
ws = W['6_기대효과']
setup(ws, '기대효과 · 3단계를 이으면 무엇이 얼마나 바뀌나', '시나리오와 핵심 값 몇 개를 바꾸면 3,000번 모의실험이 다시 돌아갑니다. 숫자는 시행 전 조건부 예상이며, 가장 큰 가정은 현장·시행 초기에 실측해 바꿉니다')
ref_sc, _ = control(ws, 7, 2, '시나리오 ▼', '확대안', options='"기본안,확대안,직접 입력"')
ref_n1, _ = control(ws, 7, 5, '참여 식당 수(직접 입력 시)', 40, '0"곳"')
ref_T, _ = control(ws, 7, 8, '운영 개월', 8, '0"개월"')
ref_n3, _ = control(ws, 7, 11, '팝업 부스 수', 8, '0"개"')
ref_pu, _ = control(ws, 7, 14, '참여율(직접 입력 시)', 0.02, '0.0%')
note(ws, 8, '시나리오: 기본안 = 식당 20곳, 참여율 실측 기반(0.1~0.9%, 안동 주민증·반값여행·하회마을) / 확대안 = 83곳, 계산대 직접 권유(0.46~5%, 가정) / 직접 입력 = 입력한 식당 수, 참여율을 중앙값으로 ±50%. '
     '나머지 가정(체험 결제율, 완주율, 반사실 등)은 데이터로 추정한 분포로 고정.', 30, size=8.5)
dZ = data_sheet('d_시나리오')
pb, pw_ = SIM['기본안']['입력'], SIM['확대안']['입력']
rows_par = [('p', pb['p'], pw_['p']), ('qb', pb['qb'], pw_['qb']), ('boat', pb['boat'], pw_['boat']), ('qp', pb['qp'], pw_['qp']), ('k', pb['k'], pw_['k'])]
dZ.append(['파라미터', '기본_하한', '기본_최빈', '기본_상한', '확대_하한', '확대_최빈', '확대_상한', '적용_하한', '적용_최빈', '적용_상한'])
for i, (k, a, b) in enumerate(rows_par, 2):
    dZ.append([k] + list(a) + list(b))
    for j, (bc, xc) in enumerate(zip('BCD', 'EFG')):
        if k == 'p':
            dZ.cell(i, 8 + j, f'=IF({ref_sc}="확대안",{xc}{i},IF({ref_sc}="직접 입력",{ref_pu}*{[0.5, 1, 1.5][j]},{bc}{i}))')
        else:
            dZ.cell(i, 8 + j, f'=IF({ref_sc}="확대안",{xc}{i},{bc}{i})')
PR = {k: (f"'d_시나리오'!$H${i}", f"'d_시나리오'!$I${i}", f"'d_시나리오'!$J${i}") for i, (k, *_ ) in enumerate(rows_par, 2)}
iN1 = len(rows_par) + 2
dZ.cell(iN1, 1, 'n1'); dZ.cell(iN1, 3, SIM['기본안']['입력']['n1']); dZ.cell(iN1, 6, SIM['확대안']['입력']['n1'])
dZ.cell(iN1, 9, f'=IF({ref_sc}="확대안",F{iN1},IF({ref_sc}="직접 입력",{ref_n1},C{iN1}))')
N1 = f"'d_시나리오'!$I${iN1}"
ws.merge_cells('B9:S9')
c9 = ws['B9']
c9.value = (f'="지금 적용된 값: 참여 식당 "&{N1}&"곳 · 참여율 "&TEXT({PR["p"][0]},"0.00%")&" ~ "&TEXT({PR["p"][2]},"0.00%")&"(최빈 "&TEXT({PR["p"][1]},"0.00%")&")'
            f' · 운영 "&{ref_T}&"개월 · 팝업 부스 "&{ref_n3}&"개"')
c9.font = F(10, True, 'BF9000'); c9.alignment = Alignment(horizontal='left', vertical='center', indent=1)
dF = data_sheet('d_분포')
EXP_P, POP, CV = P['D_체험가격']['값'], P['D_축제객단가']['값'], P['B_c']['값']
dF.append(['체험가격', '팝업객단가', '완주율'])
for i in range(max(len(EXP_P), len(POP))):
    dF.append([EXP_P[i] if i < len(EXP_P) else None, POP[i] if i < len(POP) else None, CV[i] if i < len(CV) else None])
G = SIM['고정값']
dK2 = data_sheet('d_고정')
fix = [('중구동 월 외지인 방문', G['중구동_월방문']), ('원도심 관광 식당', 83), ('월영교 주말 저녁 방문기회(8개월)', G['월영교_저녁기회']),
       ('강남동 외지인 관광소비(8개월)', G['강남동소비']), ('안동 외지인 방문(8개월)', Q['방문']['2026_1-8']), ('현재 방문당 체험·문화', Q['체험문화']['안동_2026']),
       ('목표(전국 중앙 수준)', G['목표']), ('목표까지 추가 소비(8개월)', G['격차_8개월']), ('월영야행 부스(기준)', 24), ('축제 식음 비중', 0.606),
       ('주말 일수(8개월)', 70), ('주민증 월 이용', 887.5)]
for k, v in fix: dK2.append([k, v])
FX = {k: f"'d_고정'!$B${i}" for i, (k, _) in enumerate(fix, 1)}
NS = 3000
dMC = data_sheet('d_모의실험')
mc_cols = ['u_p', 'p', 'r', '체험가격', 'c', 'u_qb', 'q_b', 'u_boat', '문보트가격', 'u_qp', 'q_p', 'u_k', 'k', '팝업객단가', '반사실',
           '인증', '낮체험', '저녁도착', '문보트', '체험결제', '체험추가', '팝업추가', '연결추가', '추가합', '방문당증가율', '격차비율']
dMC.append(mc_cols)
MCc = {n: CL(i + 1) for i, n in enumerate(mc_cols)}
def tri(u, k):
    a, m, b = PR[k]
    return f'IF({b}={a},{m},IF({u}<({m}-{a})/({b}-{a}),{a}+SQRT({u}*({b}-{a})*({m}-{a})),{b}-SQRT((1-{u})*({b}-{a})*({b}-{m}))))'
nE_, nP_, nC_ = len(EXP_P), len(POP), len(CV)
for i in range(2, NS + 2):
    c = {n: f'{MCc[n]}{i}' for n in mc_cols}
    f = {'u_p': '=RAND()', 'p': '=' + tri(c['u_p'], 'p'), 'r': '=BETAINV(RAND(),2,6)',
         '체험가격': f"=INDEX('d_분포'!$A$2:$A${nE_ + 1},RANDBETWEEN(1,{nE_}))",
         'c': f"=INDEX('d_분포'!$C$2:$C${nC_ + 1},RANDBETWEEN(1,{nC_}))*(0.9+0.2*RAND())",
         'u_qb': '=RAND()', 'q_b': '=' + tri(c['u_qb'], 'qb'), 'u_boat': '=RAND()', '문보트가격': '=' + tri(c['u_boat'], 'boat'),
         'u_qp': '=RAND()', 'q_p': '=' + tri(c['u_qp'], 'qp'), 'u_k': '=RAND()', 'k': '=' + tri(c['u_k'], 'k'),
         '팝업객단가': f"=INDEX('d_분포'!$B$2:$B${nP_ + 1},RANDBETWEEN(1,{nP_}))", '반사실': '=BETAINV(RAND(),5,5)',
         '인증': f"={FX['중구동 월 외지인 방문']}*{ref_T}*({N1}/{FX['원도심 관광 식당']})*{c['p']}",
         '낮체험': f"={c['인증']}*{c['r']}", '저녁도착': f"={c['인증']}*{c['c']}", '문보트': f"={c['저녁도착']}*{c['q_b']}",
         '체험결제': f"={c['낮체험']}+{c['문보트']}",
         '체험추가': f"=({c['낮체험']}*{c['체험가격']}+{c['문보트']}*{c['문보트가격']})*(1-{c['반사실']})",
         '팝업추가': f"=({c['저녁도착']}*{c['q_p']}*{c['팝업객단가']}+{FX['월영교 주말 저녁 방문기회(8개월)']}*({ref_T}/8)*{c['팝업객단가']}*{FX['축제 식음 비중']}*{c['k']}*({ref_n3}/{FX['월영야행 부스(기준)']}))*(1-{c['반사실']})",
         '연결추가': f"=({c['문보트']}*{c['문보트가격']}+{c['저녁도착']}*MAX({c['q_p']}-{FX['축제 식음 비중']}*{c['k']}*({ref_n3}/{FX['월영야행 부스(기준)']}),0)*{c['팝업객단가']})*(1-{c['반사실']})",
         '추가합': f"={c['체험추가']}+{c['팝업추가']}",
         '방문당증가율': f"={c['체험추가']}/({FX['안동 외지인 방문(8개월)']}*{ref_T}/8)/{FX['현재 방문당 체험·문화']}",
         '격차비율': f"={c['체험추가']}/({FX['목표까지 추가 소비(8개월)']}*{ref_T}/8)"}
    for n in mc_cols:
        dMC[c[n]] = f[n]
MR = lambda n: f"'d_모의실험'!${MCc[n]}$2:${MCc[n]}${NS + 1}"
MED = lambda n: f"MEDIAN({MR(n)})"
RANGE = lambda n, fmt: f'="90% 범위 "&TEXT(PERCENTILE({MR(n)},0.05),"{fmt}")&" ~ "&TEXT(PERCENTILE({MR(n)},0.95),"{fmt}")'
kpi(ws, 10, 2, '체험 결제(낮 + 저녁 문보트)', f"={MED('체험결제')}", '#,##0"건"', RANGE('체험결제', '#,##0'), color=RED)
kpi(ws, 10, 6, '방문당 체험·문화 소비 증가(안동 전체)', f"={MED('방문당증가율')}", '+0.0%', f'="회복 목표 +17.7% 중 "&TEXT({MED("격차비율")},"0%")&" 메움"', color=RED)
kpi(ws, 10, 10, '안동 내 추가 소비(원래 썼을 돈 제외)', f"={MED('추가합')}/100000000", '0.00"억 원"', f'="90% 범위 "&TEXT(PERCENTILE({MR("추가합")},0.05)/1E8,"0.00")&" ~ "&TEXT(PERCENTILE({MR("추가합")},0.95)/1E8,"0.00")&"억"')
kpi(ws, 10, 14, '단계를 이어야만 생기는 몫', f"={MED('연결추가')}/{MED('추가합')}", '0%', f'="저녁 문보트 결제 "&TEXT({MED("문보트")},"#,##0")&"건 (단계를 따로 하면 0)"')
kpi(ws, 15, 2, '릴레이 인증(월)', f"={MED('인증')}/{ref_T}", '#,##0"건"', f'="현재 주민증 월 이용 888건의 "&TEXT({MED("인증")}/{ref_T}/888,"0.0")&"배"')
kpi(ws, 15, 6, '저녁 원도심→월영교 이동(주말 하루)', f"={MED('저녁도착')}/(70*{ref_T}/8)", '0.0"명"', f'="4인 택시 하루 "&ROUNDUP({MED("저녁도착")}/(70*{ref_T}/8)/4,0)&"회 필요 (재배치 한도 12회)"')
kpi(ws, 15, 10, '월영교 21시 식사 가능 / 주점', f'="3→"&(3+{ref_n3})&"곳 / 0→"&{ref_n3}&"곳"', '@', '팝업 부스를 21시 이후 운영할 때')
kpi(ws, 15, 14, '회복 목표(179.7원) 도달 확률', f"=COUNTIF({MR('방문당증가율')},\">=\"&({FX['목표(전국 중앙 수준)']}/{FX['현재 방문당 체험·문화']}-1))/{NS}", '0%', '3,000번 중 목표 이상이 나온 비율')
# 분포 히스토그램
dHg = data_sheet('d_히스토')
dHg.append(['구간(방문당 증가율 %)', '회차'])
for j in range(1, 21):
    dHg.cell(j + 1, 1, f"=TEXT(PERCENTILE({MR('방문당증가율')},0.98)*{j}/20*100,\"0.0\")")
    lo = f"PERCENTILE({MR('방문당증가율')},0.98)*{j - 1}/20"; hi = f"PERCENTILE({MR('방문당증가율')},0.98)*{j}/20"
    dHg.cell(j + 1, 2, f"=COUNTIFS({MR('방문당증가율')},\">\"&{lo},{MR('방문당증가율')},\"<=\"&{hi})" if j > 1 else f"=COUNTIF({MR('방문당증가율')},\"<=\"&{hi})")
ch = BarChart(); ch.type = 'col'; ch.gapWidth = 5
s = Series(Reference(dHg, min_col=2, min_row=2, max_row=21), title='회차'); color_series(s, RED); ch.series.append(s)
ch.set_categories(Reference(dHg, min_col=1, min_row=2, max_row=21)); ch.x_axis.title = '방문당 체험·문화 소비 증가율(%, 구간 상한)'
style_chart(ch, '3,000번 모의실험 결과 분포: 방문당 체험·문화 소비 증가율', 17, 8.5, None)
ws.add_chart(ch, 'B20')
# 민감도(상관²)
dSe = data_sheet('d_민감도')
dSe.append(['값', '비중'])
sens = [('참여율', 'p'), ('낮 체험 결제율', 'r'), ('체험 가격', '체험가격'), ('자연 완주율', 'c'), ('문보트 결제율', 'q_b'), ('팝업 구매율', 'q_p'), ('팝업 강도', 'k'), ('팝업 객단가', '팝업객단가'), ('원래 썼을 돈 비율', '반사실')]
for i, (lab, k) in enumerate(sens, 2):
    dSe.cell(i, 1, lab); dSe.cell(i, 3, f"=CORREL({MR(k)},{MR('추가합')})^2")
    dSe.cell(i, 2, f"=C{i}/SUM($C$2:$C${len(sens) + 1})")
ch = BarChart(); ch.type = 'bar'
s = Series(Reference(dSe, min_col=2, min_row=2, max_row=len(sens) + 1), title='결과를 흔드는 비중'); color_series(s, NAVY); ch.series.append(s)
ch.set_categories(Reference(dSe, min_col=1, min_row=2, max_row=len(sens) + 1)); ch.x_axis.scaling.orientation = 'maxMin'; ch.y_axis.number_format = '0%'
style_chart(ch, '추가 소비를 가장 크게 흔드는 값(현장·시행 초기에 먼저 잴 값)', 17, 8.5, None)
ws.add_chart(ch, 'K20')
note(ws, 38, '계산 방법: 인증 = 중구동 외지인 방문 × (참여 식당 ÷ 83곳) × 참여율 → 낮 체험 결제·저녁 월영교 도착(자연 완주율 15~17%, 철도공사 역산) → 문보트·팝업 결제. '
     '원래 썼을 돈(반사실, 평균 50%)을 뺀 것이 추가 소비다. 참여율 하한은 주민증 26곳 이용 회귀, 최빈은 반값여행 1차 신청, 상한은 하회마을 이용률. 체험 가격(27개)·축제 객단가(56개)는 실측 분포.', 44)
note(ws, 39, '주의: 시행 전 조건부 예상이다(90% 범위 = 가정한 분포에서 나온 모의실험 구간, 신뢰구간 아님). 효과가 원도심에서 생기므로 안동 전체 %로는 작게 보인다. 사업 규모 지표(체험 결제·인증·운행)와 함께 읽는다.', 30, size=8.5)

# ═════════════ 7 사후 검증 ═════════════
ws = W['7_사후검증']
setup(ws, '사후 검증 · 효과는 원인이 생기는 곳에서 재고, 증상으로 환산한다', '도시 전체 지표로는 효과가 +24% 이상이어야 판정된다. 그래서 참여 식당을 추첨 순서로 열어 같은 달끼리 비교한다. 예상 효과를 넣어 판정 가능성을 확인해 보세요')
ref_ef, _ = control(ws, 7, 2, '안동 전체 예상 효과', 0.04, '+0.0%')
ref_de, _ = control(ws, 7, 5, '순차 확대 설계 ▼', '확대안 83곳·5개월', options='"확대안 83곳·5개월,기본안 20곳·10개월,기본안 20곳·5개월"', width=3)
ref_pe, _ = control(ws, 7, 10, '참여 식당 손님 체험 결제 증가', 1.0, '+0%')
A_ = SV['A_주민증후보']
dV = data_sheet('d_검증')
dV.append(['분기', '안동', '비교 조합'])
for qn, a, b in zip(SV['시계열']['분기'], SV['시계열']['안동'], SV['시계열']['합성']):
    dV.append([f"{qn[2:4]}.{qn[4]}Q", a, b])
nV = len(SV['시계열']['분기']) + 1
dV.cell(1, 5, '효과 δ'); dV.cell(1, 6, '한 방향 p(주민증 후보)'); dV.cell(1, 7, '한 방향 p(전국 시군)')
for i, (ca, cb) in enumerate(zip(A_['곡선'], SV['B_전국시군']['곡선']), 2):
    dV.cell(i, 5, ca['δ']); dV.cell(i, 6, ca['p_한방향']); dV.cell(i, 7, cb['p_한방향'])
nCv = len(A_['곡선']) + 1
dV.cell(1, 9, '효과'); dV.cell(1, 10, '확대안 83곳·5개월'); dV.cell(1, 11, '기본안 20곳·10개월'); dV.cell(1, 12, '기본안 20곳·5개월')
effs = [c['효과'] for c in PW['확대안']['곡선']]
for i, e in enumerate([0.0] + effs, 2):
    dV.cell(i, 9, e)
    for j, key in enumerate(('확대안', '기본안_2개월', '기본안'), 10):
        val = PW[key]['효과0_거짓양성'] if e == 0 else next(c['검정력'] for c in PW[key]['곡선'] if abs(c['효과'] - e) < 1e-9)
        dV.cell(i, j, val)
nPw = len(effs) + 2
mde = A_['MDE_한방향_p10']
kpi(ws, 9, 2, '도시 합성통제로 판정 가능?', f'=IF({ref_ef}>={mde},"가능","어려움")', '@', f'최소 검출 효과 +{mde * 100:.0f}% (한 방향, p ≤ 0.10)', color=RED)
pe_ = f"MIN(MAX({ref_pe},0),{max(effs)})"
m_ = f"MATCH({pe_},'d_검증'!$I$2:$I${nPw},1)"
cc_ = f"MATCH({ref_de},'d_검증'!$J$1:$L$1,0)"
y0 = f"INDEX('d_검증'!$J$2:$L${nPw},{m_},{cc_})"; y1 = f"INDEX('d_검증'!$J$2:$L${nPw},MIN({m_}+1,{nPw - 1}),{cc_})"
x0 = f"INDEX('d_검증'!$I$2:$I${nPw},{m_})"; x1 = f"INDEX('d_검증'!$I$2:$I${nPw},MIN({m_}+1,{nPw - 1}))"
kpi(ws, 9, 6, '순차 확대 검정력', f"=IFERROR({y0}+({pe_}-{x0})*({y1}-{y0})/({x1}-{x0}),{y0})", '0%', '입력한 효과를 판정할 확률(80% 이상이면 충분)', color='548235')
kpi(ws, 9, 10, '안동 사전 적합 순위', A_['사전RMSPE_순위(낮을수록 좋음)'], '0"위"', '후보 49곳 중(비교 조합이 과거를 가장 잘 따라감)')
kpi(ws, 9, 14, '효과 없는 기간의 거짓 양성', A_['가짜시행_p_한방향'], '0.00', '가짜 시행 한 방향 p (높을수록 정상)')
ch = LineChart()
for col, colr in ((2, RED), (3, GRAY)):
    s = Series(Reference(dV, min_col=col, min_row=2, max_row=nV), title=dV.cell(1, col).value); color_series(s, colr, True); ch.series.append(s)
ch.set_categories(Reference(dV, min_col=1, min_row=2, max_row=nV)); ch.y_axis.title = '지수(2022~2023 평균=100)'
style_chart(ch, '안동 vs 비교 조합(태백·영동·거창·영덕 등): 2025.2Q까지 사전, 이후 가짜 시행 구간', 17, 8.5)
ws.add_chart(ch, 'B14')
ch = LineChart()
for col, colr in ((6, RED), (7, GRAY)):
    s = Series(Reference(dV, min_col=col, min_row=2, max_row=nCv), title=dV.cell(1, col).value); color_series(s, colr, True); ch.series.append(s)
ch.set_categories(Reference(dV, min_col=5, min_row=2, max_row=nCv)); ch.x_axis.number_format = '0%'; ch.x_axis.tickLblSkip = 10; ch.y_axis.title = 'p값'
style_chart(ch, '도시 합성통제: 효과 크기 δ별 p값(0.10 아래로 내려가야 판정)', 17, 8.5)
ws.add_chart(ch, 'K14')
ch = LineChart()
for col, colr in ((10, '548235'), (11, BLUE), (12, GRAY)):
    s = Series(Reference(dV, min_col=col, min_row=2, max_row=nPw), title=dV.cell(1, col).value); color_series(s, colr, True); ch.series.append(s)
ch.set_categories(Reference(dV, min_col=9, min_row=2, max_row=nPw)); ch.x_axis.number_format = '0%'; ch.y_axis.number_format = '0%'; ch.y_axis.title = '검정력'
style_chart(ch, '순차 확대(추첨): 참여 식당 손님 체험 결제 증가별 검정력(모의실험 400회)', 34, 8.5)
ws.add_chart(ch, 'B32')
note(ws, 50, '읽는 법: 도시 지표(안동 전체)는 효과가 1,172만 방문에 희석되고 분기마다 크게 출렁여 +24% 이상만 판정된다. 참여 식당 손님 단위에서는 같은 효과가 약 2배(+100%, 반사실 50%)로 보여, '
     '추첨 순서로 연 묶음과 아직 안 연 묶음을 같은 달에 비교하면 판정할 수 있다. 확인된 효과는 다시 "체험·문화 격차 중 몇 %를 메웠나"로 환산해 보고한다. 과거에 비교 조합이 안동을 잘 따라온 것은 방법이 쓸 만하다는 뜻이지 정책 효과의 증명이 아니다.', 50)
note(ws, 51, '출처: 데이터랩 239개 시·군 월별 외지인 카드·방문 패널(2022.1~2026.6, 2024.9~12 카드 결측), 합성통제(사전 추세 상관 상위 10곳, 한 방향 순위 검정), 순차 확대 검정력은 포아송 고정효과 모의실험', 20, size=8.5, color='7F7F7F')

# ═════════════ 8 전국 진단 ═════════════
ws = W['8_전국진단']
setup(ws, '전국 진단 · 안동과 같은 문제를 가진 시·군 찾기', '시·군을 고르면 네 지표에서의 위치와 문제 유형이 나온다. 가중치를 바꿔 순위가 얼마나 흔들리는지도 확인해 보세요(점수보다 유형으로 읽는다)')
dDg = data_sheet('d_진단')
nD = len(dg) + 1
dDg.append(['시군'] + IND + ['문제①', '문제②', '문제③', '문제④', '점수', '순위', '유형'])
for i, (_, r) in enumerate(dg.iterrows(), 2):
    dDg.cell(i, 1, r['시군'])
    for j, c in enumerate(IND):
        dDg.cell(i, 2 + j, float(r[c]))
ref_sel, _ = control(ws, 7, 2, '시·군 ▼', '안동시', options=f"='d_진단'!$A$2:$A${nD}")
wref = []
for j, lab in enumerate(['가중치: 체험', '숙박', '저녁', '야간']):
    rr, _ = control(ws, 7, 5 + j * 3, lab, 1, '0.0', width=2)
    wref.append(rr)
for i in range(2, nD + 1):
    for j in range(4):
        col = CL(2 + j)
        dDg.cell(i, 6 + j, f"=100*(1-(COUNTIF(${col}$2:${col}${nD},\"<\"&{col}{i})+1)/COUNT(${col}$2:${col}${nD}))")
    dDg.cell(i, 10, f"=IF(({wref[0]}+{wref[1]}+{wref[2]}+{wref[3]})<=0,AVERAGE(F{i}:I{i}),(F{i}*{wref[0]}+G{i}*{wref[1]}+H{i}*{wref[2]}+I{i}*{wref[3]})/({wref[0]}+{wref[1]}+{wref[2]}+{wref[3]}))")
    dDg.cell(i, 11, f"=RANK(J{i},$J$2:$J${nD},0)")
    dDg.cell(i, 12, f'=TRIM(IF(F{i}>=75,"체험 ","")&IF(G{i}>=75,"숙박 ","")&IF(H{i}>=75,"저녁 ","")&IF(I{i}>=75,"야간",""))')
L = lambda col: f"INDEX('d_진단'!${col}$2:${col}${nD},MATCH({ref_sel},'d_진단'!$A$2:$A${nD},0))"
kpi(ws, 9, 2, '진단 점수(0=좋음, 100=나쁨)', f"={L('J')}", '0.0', f'="{len(dg)}곳 중 "&{L("K")}&"위"', color=RED)
kpi(ws, 9, 6, '문제 유형(하위 25% 지표)', f'=IF({L("L")}="","해당 없음",{L("L")})', '@', f'="같은 유형 "&COUNTIF(\'d_진단\'!$L$2:$L${nD},{L("L")})&"곳"')
kpi(ws, 9, 10, '체험·문화 소비 변화율', f"={L('B')}", '+0.0%;-0.0%', f'="문제 백분위 "&TEXT({L("F")},"0")&" · 중앙 "&TEXT(MEDIAN(\'d_진단\'!$B$2:$B${nD}),"+0.0%;-0.0%")')
kpi(ws, 9, 14, '평균 숙박일수', f"={L('C')}", '0.00"일"', f'="문제 백분위 "&TEXT({L("G")},"0")&" · 중앙 "&TEXT(MEDIAN(\'d_진단\'!$C$2:$C${nD}),"0.00")&"일"')
dPc = data_sheet('d_진단비교')
dPc.append(['지표', '선택 시·군', '안동'])
for j, lab in enumerate(['체험·문화 변화', '숙박일수', '저녁 전환율', '야간 방문 비중']):
    col = CL(6 + j)
    dPc.append([lab, f"={L(col)}", f"=INDEX('d_진단'!${col}$2:${col}${nD},MATCH(\"안동시\",'d_진단'!$A$2:$A${nD},0))"])
ch = BarChart(); ch.type = 'col'
for col, colr in ((2, BLUE), (3, RED)):
    s = Series(Reference(dPc, min_col=col, min_row=2, max_row=5), title=dPc.cell(1, col).value); color_series(s, colr); ch.series.append(s)
ch.set_categories(Reference(dPc, min_col=1, min_row=2, max_row=5)); ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 100; ch.y_axis.title = '문제 백분위'
style_chart(ch, '네 지표의 문제 백분위: 선택 시·군 vs 안동 (75 이상 = 하위 25%)', 17, 9)
ws.add_chart(ch, 'B14')
sc = ScatterChart(); sc.style = 13
dPc.cell(7, 1, '점'); dPc.cell(7, 2, 'x(체험 변화율)'); dPc.cell(7, 3, 'y(숙박일수)')
dPc.cell(8, 1, '선택'); dPc.cell(8, 2, f"={L('B')}"); dPc.cell(8, 3, f"={L('C')}")
dPc.cell(9, 1, '안동'); dPc.cell(9, 2, f"=INDEX('d_진단'!$B$2:$B${nD},MATCH(\"안동시\",'d_진단'!$A$2:$A${nD},0))"); dPc.cell(9, 3, f"=INDEX('d_진단'!$C$2:$C${nD},MATCH(\"안동시\",'d_진단'!$A$2:$A${nD},0))")
for (sh, xc, yc, r1, r2, t, colr, sz) in ((dDg, 2, 3, 2, nD, '시·군', 'C9C9C9', 4), (dPc, 2, 3, 8, 8, '선택', BLUE, 10), (dPc, 2, 3, 9, 9, '안동', RED, 10)):
    s_ = Series(Reference(sh, min_col=yc, min_row=r1, max_row=r2), Reference(sh, min_col=xc, min_row=r1, max_row=r2), title=t)
    s_.marker.symbol = 'circle'; s_.marker.size = sz; s_.marker.graphicalProperties.solidFill = colr; s_.marker.graphicalProperties.line.solidFill = colr
    s_.graphicalProperties.line.noFill = True; sc.series.append(s_)
sc.x_axis.title = '체험·문화 소비 변화율'; sc.y_axis.title = '평균 숙박일수'; sc.x_axis.number_format = '0%'
sc.x_axis.scaling.min = -0.6; sc.x_axis.scaling.max = 1.0; sc.y_axis.scaling.min = 2; sc.y_axis.scaling.max = 5.5
style_chart(sc, '체험·문화 소비 변화 × 숙박일수: 왼쪽 아래 = 안동과 같은 유형', 17, 9)
ws.add_chart(sc, 'K14')
label(ws, 'B33', '진단 점수 상위 10곳(가중치 반영, 실시간)', 10.5, True, NAVY)
hdr = ['순위', '시·군', '점수', '유형']
SPAN = [(2, 3), (4, 5), (6, 7), (8, 11)]
for r in range(34, 45):
    for c1, c2 in SPAN:
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        for cc in range(c1, c2 + 1): ws.cell(r, cc).border = BOX
for j, h in enumerate(hdr):
    c = ws.cell(34, SPAN[j][0], h); c.font = F(9, True); c.fill = FILL('D9E1F2'); c.alignment = CENTER
for k in range(1, 11):
    r = 34 + k
    ws.cell(r, 2, k)
    ws.cell(r, 4, f"=INDEX('d_진단'!$A$2:$A${nD},MATCH(LARGE('d_진단'!$J$2:$J${nD},{k}),'d_진단'!$J$2:$J${nD},0))")
    ws.cell(r, 6, f"=LARGE('d_진단'!$J$2:$J${nD},{k})").number_format = '0.0'
    ws.cell(r, 8, f"=INDEX('d_진단'!$L$2:$L${nD},MATCH(LARGE('d_진단'!$J$2:$J${nD},{k}),'d_진단'!$J$2:$J${nD},0))")
    for cc in (2, 4, 6, 8):
        ws.cell(r, cc).font = F(9.5); ws.cell(r, cc).alignment = CENTER
same = dg[dg['유형'] == dg.loc[dg['시군'] == '안동시', '유형'].values[0]]['시군'].tolist()
note(ws, 46, f'안동과 같은 유형(체험·문화 소비 하락 + 짧은 숙박) {len(same)}곳: {", ".join(same)} → 이어드림 방식의 확산 후보. '
     '같은 가중치로 안동은 57위지만 가중치를 무작위로 바꾸면 18~100위로 흔들린다 → 점수 하나로 순위를 매기지 않고 유형으로 읽는다.', 34)
note(ws, 47, '지표(낮을수록 문제): ① 방문당 체험·문화 소비 변화율 2024→2026 1~8월 ② 평균 숙박일수 2025(LN_02_01_013) ③ 저녁 전환율 18~21시÷14~18시 ④ 야간(21~24시) 방문 비중, 2026 1~8월(BDT_01_01_006). 시·군 144곳(방문 100만 이상, 광역시 자치구 제외). 탐색 지표이며 정책 효과 점수가 아니다.', 30, size=8.5, color='7F7F7F')

# ═════════════ 9 서식4 ═════════════
ws = W['9_서식4']
setup(ws, '서식4 활용사례 작성본(초안)', '공모요강 서식4 양식 · 개조식 2~3장 · 함초롬바탕 11pt · 줄간격 160. 전체 파일은 보고서/서식4_20260923/')
txt = (ROOT / '보고서/서식4_20260923/서식4_미리보기.html').read_text(encoding='utf-8')
body = re.search(r'<body>(.*)</body>', txt, re.S).group(1)
body = re.sub(r'<div class="fig">.*?</div>', '', body, flags=re.S)
body = re.sub(r'<tr>', '\n', body); body = re.sub(r'</t[dh]>', ' | ', body)
body = re.sub(r'<(div|p|h1)[^>]*>', '\n', body); body = re.sub(r'<[^>]+>', '', body)
lines = [H.unescape(l).strip() for l in body.split('\n') if H.unescape(l).strip()]
for i, l in enumerate(lines, 7):
    ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=19)
    c = ws.cell(i, 2, l)
    head = re.match(r'^\d\) ', l) is not None
    c.font = F(10.5 if head else 10, head, NAVY if head else INK); c.alignment = LEFT
    if head:
        for cc in range(2, 20): ws.cell(i, cc).fill = FILL('DDEBF7')
    ws.row_dimensions[i].height = max(18, 15 * (len(l) // 95 + 1))

# ═════════════ 10 방문객 설문 ═════════════
ws = W['10_설문']
svn = SVY['표본']
setup(ws, '방문객 설문 · 방문객이 직접 말한 불편과 의향', f"팀 온라인 설문 {svn['응답']}명(최근 3년 안에 안동 방문 {svn['3년내_방문_예']}명, 2026.9.23~24). 응답자 범위와 이동수단을 바꿔 보세요. "
      '불편·부족은 근거로, 의향은 참여율이 아니라 수요의 방향으로만 읽습니다')
ref_sp, _ = control(ws, 7, 2, '응답자 범위 ▼', '방문자만', options='"방문자만,전체"')
ref_md, _ = control(ws, 7, 5, '이동수단 ▼', '전체', options='"전체,자가용·렌터카,차 없음"')
sv = pd.read_csv(ROOT / 'data/설문/안동관광온라인설문_응답_정제.csv')
LK = {'매우있다': 5, '매우 있다': 5, '있다': 4, '보통이다': 3, '별로없다': 2, '별로 없다': 2, '별로없다(자차/렌터카 이용 등)': 2, '전혀없다': 1, '전혀 없다': 1}
q2 = sv['Q2_주_이동수단'].fillna('')
q4s = sv['Q4_이동불편_포기한곳'].fillna('')
K11 = {'체험 정보가 부족해서': '정보 부족', '줄을 오래 서거나 예약하기 불편해서': '줄·예약 불편', '다음 장소로 이동할 시간이 부족해서': '이동 시간 부족', '가격이 부담스러워서': '가격 부담'}
K13 = {'18시 이전': '18시 이전', '18시 ~ 20시': '18-20시', '20시 ~ 22시 (막차 시간대)': '20-22시(막차)', '당일 귀가 안 함 (안동에서 숙박)': '숙박'}
dS = data_sheet('d_설문')
dS.append(['방문', '이동', '버스주이용', 'Q3', '포기_외곽', '포기_월영교', '포기_원도심', '포기_없음', 'Q4응답', 'Q9', 'Q11', 'Q13', 'Q6', 'Q8', 'Q10', 'Q12'])
for _, r in sv.iterrows():
    mode = '자가용' if q2[_] == '자가용 / 렌터카' else ('미응답' if q2[_] == '' else '차 없음')
    q9v = r['Q9_월영교_밤_식당편의']
    q9c = '부족' if q9v in ('부족했다', '매우부족했다') else ('보통' if q9v == '보통이다' else ('충분' if q9v == '충분하다' else '미응답'))
    q11 = r['Q11_유료체험_망설임']; q11c = K11.get(q11, '미응답' if pd.isna(q11) else '기타')
    dS.append([r['Q1_3년내_방문'], mode, int(q2[_] == '대중교통 (시내버스 등)'), r['Q3_대중교통_불편'] if pd.notna(r['Q3_대중교통_불편']) else '미응답',
               int('외곽 관광지' in q4s[_]), int('야간 명소' in q4s[_]), int('원도심의 밤' in q4s[_]), int(q4s[_].startswith('없음')), int(q4s[_] != ''),
               q9c, q11c, K13.get(r['Q13_귀가_교통_시간대'], '미응답')] +
              [LK.get(r[c], 0) for c in ('Q6_영수증_체험쿠폰_의향', 'Q8_야간택시셔틀_의향', 'Q10_야간팝업포차_의향', 'Q12_영수증_축제혜택_의향')])
nS = len(sv) + 1
rg = lambda col: f"'d_설문'!${col}$2:${col}${nS}"
dSc = data_sheet('d_설문계산')
dSc['A1'] = '방문 조건'; dSc['B1'] = f'=IF({ref_sp}="방문자만","예","*")'
dSc['A2'] = '이동 조건'; dSc['B2'] = f'=IF({ref_md}="전체","*",IF({ref_md}="자가용·렌터카","자가용","차 없음"))'
BASE = f"{rg('A')},'d_설문계산'!$B$1,{rg('B')},'d_설문계산'!$B$2"
cnt = lambda extra='': f"COUNTIFS({BASE}{',' + extra if extra else ''})"
dSc['A3'] = '선택 응답자 수'; dSc['B3'] = '=' + cnt()
note(ws, 7, "=\"선택한 응답자 \"&'d_설문계산'!$B$3&\"명\"", col1=9, col2=13, size=10, color=NAVY)
def pct_card(col, lab, num, den, sub_txt, color=NAVY):
    f_ = f"=IFERROR({cnt(num)}/{cnt(den)},\"해당 없음\")"
    kpi(ws, 9, col, lab, f_, '0%', f"=\"{sub_txt} \"&{cnt(num)}&\"명 / \"&{cnt(den)}&\"명\"", color=color)
pct_card(2, '시내버스 이용자 중 "불편했다"', f"{rg('C')},1,{rg('D')},\"불편했다\"", f"{rg('C')},1", '주 이동수단이 시내버스인 사람 중', RED)
pct_card(6, '이동이 불편해 월영교 등 야간 명소 포기', f"{rg('I')},1,{rg('F')},1", f"{rg('I')},1", '포기한 곳 문항 응답자 중', RED)
pct_card(10, '20시 이후 월영교 식당·편의시설 부족', f"{rg('J')},\"부족\"", f"{rg('J')},\"<>미응답\"", '부족·매우 부족', RED)
pct_card(14, '당일 귀가자 중 20~22시(막차) 귀가', f"{rg('L')},\"20-22시(막차)\"", f"{rg('L')},\"<>미응답\",{rg('L')},\"<>숙박\"", '숙박자 제외')
# 의향 4문항 5점 분포
IT = [('M', '체험 쿠폰(영수증 인증)'), ('N', '야간 택시·셔틀(18:30~21:00)'), ('O', '월영교 야간 팝업(21~23시)'), ('P', '축제 혜택(영수증)')]
LV = [(5, '매우 있다', NAVY), (4, '있다', BLUE), (3, '보통', 'BFBFBF'), (2, '별로 없다', 'F4B183'), (1, '전혀 없다', RED)]
dSc['D1'] = '문항'
for j, (_, t, _c) in enumerate(LV): dSc.cell(1, 5 + j, t)
for i, (col, t) in enumerate(IT, 2):
    dSc.cell(i, 4, t)
    for j, (v, _t, _c) in enumerate(LV):
        dSc.cell(i, 5 + j, f"=IFERROR({cnt(f'{rg(col)},{v}')}/{cnt(f'{rg(col)},\">0\"')},0)")
ch = BarChart(); ch.type = 'bar'; ch.grouping = 'percentStacked'; ch.overlap = 100; ch.gapWidth = 60
for j, (_v, t, colr) in enumerate(LV):
    s = Series(Reference(dSc, min_col=5 + j, min_row=2, max_row=5), title=t); color_series(s, colr); ch.series.append(s)
ch.set_categories(Reference(dSc, min_col=4, min_row=2, max_row=5)); ch.x_axis.scaling.orientation = 'maxMin'; ch.y_axis.number_format = '0%'
style_chart(ch, '이용 의향(5점): 선택한 응답자 중 비율, 의향은 참여율이 아니다', 34, 8)
ws.add_chart(ch, 'B14')
# 포기한 곳 · 망설임 이유
dSc['K1'] = '포기한 곳'; dSc['L1'] = '비율'
for i, (col, t) in enumerate((('E', '외곽 관광지(하회·도산)'), ('F', '야간 명소(월영교)'), ('G', '원도심 밤 가게·식당'), ('H', '없음(모두 방문)')), 2):
    dSc.cell(i, 11, t); dSc.cell(i, 12, f"=IFERROR({cnt(f'{rg(col)},1,' + rg('I') + ',1')}/{cnt(rg('I') + ',1')},0)")
ch = BarChart(); ch.type = 'bar'
s = Series(Reference(dSc, min_col=12, min_row=2, max_row=5), title='비율'); color_series(s, RED); ch.series.append(s)
ch.set_categories(Reference(dSc, min_col=11, min_row=2, max_row=5)); ch.x_axis.scaling.orientation = 'maxMin'; ch.y_axis.number_format = '0%'
style_chart(ch, '이동이 불편해 가고 싶었지만 포기한 곳(중복 선택)', 17, 8.5, None)
ws.add_chart(ch, 'B31')
dSc['N1'] = '망설임 이유'; dSc['O1'] = '비율'
for i, t in enumerate(['정보 부족', '줄·예약 불편', '이동 시간 부족', '가격 부담', '기타'], 2):
    dSc.cell(i, 14, t); dSc.cell(i, 15, f"=IFERROR({cnt(rg('K') + ',\"' + t + '\"')}/{cnt(rg('K') + ',\"<>미응답\"')},0)")
ch = BarChart(); ch.type = 'bar'
s = Series(Reference(dSc, min_col=15, min_row=2, max_row=6), title='비율'); color_series(s, NAVY); ch.series.append(s)
ch.set_categories(Reference(dSc, min_col=14, min_row=2, max_row=6)); ch.x_axis.scaling.orientation = 'maxMin'; ch.y_axis.number_format = '0%'
style_chart(ch, '유료 공연·체험이 망설여지는 가장 큰 이유: 가격보다 정보·예약·시간', 17, 8.5, None)
ws.add_chart(ch, 'K31')
wt = SVY['이동수단_가중_방문자']
note(ws, 49, '이동수단 보정: 응답한 방문자 중 차 없는 사람이 46%로 실제(철도공사 추정 철도+버스 10.5%)보다 많다. 10.5 : 89.5로 다시 가중하면 긍정(있다+매우 있다) 비율은 '
     f"체험 쿠폰 {wt['Q6_영수증_체험쿠폰_의향']['가중_top2']:.0f}% · 셔틀 {wt['Q8_야간택시셔틀_의향']['가중_top2']:.0f}%(차 없음 {wt['Q8_야간택시셔틀_의향']['차없음_top2']:.0f}%, 자가용 {wt['Q8_야간택시셔틀_의향']['자차_top2']:.0f}%) · "
     f"팝업 {wt['Q10_야간팝업포차_의향']['가중_top2']:.0f}% · 축제 혜택 {wt['Q12_영수증_축제혜택_의향']['가중_top2']:.0f}%.", 32)
note(ws, 50, '읽는 법: "불편했다·부족했다·막차에 귀가"는 방문 경험이라 기존 데이터(112번 막차 18:45, 월영교 1km 주점 0곳, 안동역 막차 21~22시)를 뒷받침하는 근거로 쓴다. '
     '"이용하겠다"는 의향이라 실제 참여율로 쓰지 않는다. 기대효과 모의실험의 체험 결제율 r(평균 25%)은 그대로 두고, "매우 있다" 비율을 상한 참고로만 본다. '
     '가격 때문에 망설인다는 응답은 14%뿐이라 1단계 혜택은 할인 폭보다 체험 정보·예약·동선을 잇는 데 무게를 둔다.', 44)
note(ws, 51, "출처: 팀 온라인 설문(구글폼, 2026-09-23~24, 응답 103명, 편의 표본, 인구통계·개인정보 문항 없음). 안 가 본 응답자도 뒤 문항에 답해 경험 문항은 '방문자만'으로 보는 것이 기본. "
     '정제 scripts/설문_정제.py · 집계 scripts/설문_분석.py', 30, size=8.5, color='7F7F7F')

# ═════════════ 11 출처 ═════════════
ws = W['11_출처']
setup(ws, '데이터 출처 · 무엇을 어디에 썼나', '한국관광 데이터랩 8종 + 타분야 7종 + 현장조사 + 팀 온라인 설문. 모든 수치는 분석 스크립트로 다시 계산할 수 있다')
src = pd.DataFrame([
    ('데이터랩', '외지인 신용카드 관광소비(업종 중분류)', 'BDT_02_01_003', '2019~2026', '문제, 사후 검증, 전국 진단'),
    ('데이터랩', '방문자 수 성별·연령·요일·시간대 교차표(외지인)', 'BDT_01_01_006', '2019~2026', '문제, 기대효과, 전국 진단'),
    ('데이터랩', '읍면동별 외지인 방문자 수', 'BDT_01_01_005_1', '2018~2026', '혜택 배치, 기대효과'),
    ('데이터랩', '평균 숙박일수', 'LN_02_01_013', '2025', '전국 진단'),
    ('데이터랩', '관심관광지·목적지 검색(내비)', 'LN_03_01_037 · BDT_03_01_003_1', '2018~2026', '문제(하회마을 검색 비중)'),
    ('데이터랩', '야간관광 현황(야간 검색·방문)', 'BY_TH_NIGHT_TOUR', '2024~2026', '월영교, 기대효과(완주율)'),
    ('데이터랩', '축제 방문자·관광소비', 'FE_01_01_005 · 007', '2025', '기대효과(팝업 객단가)'),
    ('데이터랩', '지역별 관광지출액(강남동 외지인)', '화면 전사', '2023~2026', '기대효과, 월영야행 사건연구'),
    ('타분야', '안동역 시간대별 승하차(이음·새마을·무궁화)', '코레일(팀원 취합)', '2024.1~2026.8', '운영 시간'),
    ('타분야', '8대 도시(안동) 관광 형태 분석(가명결합)', '한국철도공사', '2022.4~6', '월영교, 기대효과(완주율)'),
    ('타분야', '상가(상권)정보 음식점 3,254곳', '소상공인시장진흥공단', '2026.6', '혜택 배치, 월영교'),
    ('타분야', '디지털관광주민증 혜택업체·이용 건수', '안동시 정보공개청구', '2024.6~2026.8', '혜택 배치, 기대효과'),
    ('타분야', '주민증 가맹점별 이용 표시', '한국관광공사 주민증 누리집', '2026-09-21', '혜택 배치, 기대효과(회귀)'),
    ('타분야', '주요관광지점 입장객', '관광지식정보시스템', '2023~2026.6', '기대효과(하회마을 이용률)'),
    ('타분야', '시내버스 시간표·노선·정류장', '안동시 버스정보시스템', '2026-09-22', '버스(BIS)'),
    ('현장', '월영교 1km 음식점 영업 종료', '팀 현장조사', '2026-09', '월영교'),
    ('설문', '팀 온라인 설문 103명(방문 경험 72명)', '구글폼(팀 직접)', '2026-09-23~24', '방문객 설문, 혜택 배치, 운영 시간, 버스, 월영교'),
    ('전화', '관광택시 운행 관행 · 월영야행 부지', '안동시관광협의회 · 한국정신문화재단', '2026-09-22', '기대효과, 월영교'),
], columns=['구분', '데이터', '제공처·코드', '기간', '쓴 화면'])
put_table(ws, 7, 2, src, font_size=9.5)
for col, w_ in zip('BCDEF', (10, 40, 30, 14, 30)): ws.column_dimensions[col].width = w_
note(ws, 27, '분석 스크립트(저장소 scripts/): 교수브리핑_수치.py · 이어드림_파라미터추정.py · 이어드림_시뮬레이션.py · 이어드림_사후검증_패널.py · 이어드림_사후검증_합성통제.py · 이어드림_순차도입_검정력.py · 전국진단표.py · 설문_정제.py · 설문_분석.py · 교수님_통합대시보드.py', 34, size=9)

# 시트 순서: 화면 먼저, 데이터 시트 뒤
order = [n for n, _ in SHEETS] + [s.title for s in wb.worksheets if s.title.startswith('d_')]
wb._sheets = [wb[n] for n in order]
wb.active = 0
from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(calcMode='auto', fullCalcOnLoad=True)          # 엑셀이 열 때 모든 수식 계산(LibreOffice 재저장 없이 전달)
OUT.mkdir(parents=True, exist_ok=True)
xl = OUT / '안동이어드림_데이터대시보드.xlsx'
wb.save(xl)
print(xl.relative_to(ROOT), len(wb.sheetnames), '시트')
