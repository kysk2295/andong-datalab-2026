"""안동팀 작업 목록 1~12번 중 빠진 자료 보강 (2026-09-22).

01 시내버스 노선·시간표 / 04 BIS 정류장 위치(안내기 설치 여부 포함)를 만든다.
원자료: data/external/작업목록_1-12_보강_20260922/01_시내버스/
  - busTimetable.json          안동시 버스정보시스템 누리집 '버스노선 시간표'가 쓰는 공개 API(/api/hp/busTimetable)
  - api_hp_busRoute.json       노선 변형 목록(/api/hp/busRoute)
  - api_bus_routes_stops_all.json  노선별 정류장 순서(/api/bus/routes/stops)
  - api_bus_stops.json         전 정류장 + 안내기(BIT) 설치 여부 bitYn(/api/bus/stops)
  - (국토부) ../04_정류장/ 2025-10-31 — 팀이 받은 ~/Downloads/국토교통부_전국 버스정류장 위치정보_안동시/ 사본
규칙: 작업 목록 2장(출처 URL·게시일·조사일, 추정 금지, 계산은 수식).
"""
import datetime as dt
import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

import importlib.util

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data/external/작업목록_1-12_보강_20260922/01_시내버스"
OUT = ROOT / "조사"
TODAY = dt.date(2026, 9, 22)
BIS_TT = "https://bus.andong.go.kr/timetable"
BIS = "https://bus.andong.go.kr/"
MOLIT = "https://www.data.go.kr/data/15067528/fileData.do"  # 국토교통부_전국 버스정류장 위치정보

# 6~11번 스크립트의 서식 도우미를 그대로 쓴다(같은 모양의 파일)
_spec = importlib.util.spec_from_file_location("base", ROOT / "scripts/작업목록_6-11_엑셀생성.py")
base = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(base)
write_table, notes_sheet, FONT, BOLD = base.write_table, base.notes_sheet, base.FONT, base.BOLD

# 기준점 = 데이터랩 LN_03_01_042 POI 좌표(scripts/월영교_반경_음식점.py와 같은 값)
ANCHOR = {"안동역": (128.6748, 36.5747, 300), "원도심": (128.728, 36.5655, 500), "월영교": (128.7609, 36.5765, 300)}
PAIRS = [("안동역", "원도심"), ("원도심", "안동역"), ("원도심", "월영교"), ("월영교", "원도심"), ("안동역", "월영교"), ("월영교", "안동역")]
HTYPE = {"1": "평일", "2": "공휴일(일요일)"}


def hav(lon1, lat1, lon2, lat2):
    R, p = 6371000, np.pi / 180
    h = np.sin((lat2 - lat1) * p / 2) ** 2 + np.cos(lat1 * p) * np.cos(lat2 * p) * np.sin((lon2 - lon1) * p / 2) ** 2
    return 2 * R * np.arcsin(np.sqrt(h))


def load():
    stops = pd.DataFrame(json.load(open(RAW / "api_bus_stops.json", encoding="utf-8")))
    rs = pd.DataFrame(json.load(open(RAW / "api_bus_routes_stops_all.json", encoding="utf-8")))
    tt = pd.DataFrame(json.load(open(RAW / "busTimetable.json", encoding="utf-8")))
    rt = pd.DataFrame(json.load(open(RAW / "api_hp_busRoute.json", encoding="utf-8")))
    area = {}
    for nm, (x, y, rad) in ANCHOR.items():
        d = hav(x, y, stops.lng.values, stops.lat.values)
        stops["d_" + nm] = d
        area[nm] = set(stops.loc[d <= rad, "stopId"])
    return stops, rs, tt, rt, area


def segments(rs, area):
    """노선 변형별로 A권역 정류장 다음에 B권역 정류장이 오는지 → (routeId, 구간) 목록."""
    out = []
    for rid, g in rs.groupby("routeId"):
        g = g.sort_values("routeOrd")
        first = g.routeOrd.min()
        pos = {nm: g[g.stopId.isin(s)] for nm, s in area.items()}
        for a, b in PAIRS:
            if pos[a].empty or pos[b].empty:
                continue
            ia = pos[a].routeOrd.min()
            later = pos[b][pos[b].routeOrd > ia]
            if later.empty:
                continue
            ib = later.routeOrd.min()
            sa = pos[a].loc[pos[a].routeOrd == ia, "stopName"].iloc[0]
            sb = later.loc[later.routeOrd == ib, "stopName"].iloc[0]
            out.append({"routeId": rid, "구간": f"{a}→{b}", "시작정류장": sa, "끝정류장": sb,
                        "기점부터_정류장수": int(ia - first), "구간_정류장수": int(ib - ia)})
    return pd.DataFrame(out)


def to_time(s):
    h, m = [int(x) for x in s.strip().split(":")]
    return dt.time(h, m)


def build_01():
    stops, rs, tt, rt, area = load()
    seg = segments(rs, area)
    seg_by_rid = seg.groupby("routeId")["구간"].apply(set).to_dict()
    # 출발시각 전체(1행 = 1회 출발)
    dep = []
    for _, r in tt.iterrows():
        times = [t for t in str(r.strtTm or "").split(",") if t.strip()]
        for t in times:
            dep.append({"routeId": r.routeId, "노선번호": r.routeNum, "노선명": r.routeNm, "시간표 구분": HTYPE.get(str(r.holidayType), str(r.holidayType)),
                        "기점": r.stStationNm, "종점": r.edStationNm, "출발시각": to_time(t)})
    dep = pd.DataFrame(dep)

    wb = Workbook()
    ws = wb.active
    ws.title = "구간별_요약"
    rows = []
    for (g, num), s in seg.merge(rt[["routeId", "routeNum"]], on="routeId").groupby(["구간", "routeNum"]):
        ids = set(s.routeId)
        d = dep[dep.routeId.isin(ids)]
        wk, hol = d[d["시간표 구분"] == "평일"], d[d["시간표 구분"] != "평일"]
        if d.empty:
            continue
        names = sorted(set(d["노선명"]))
        rows.append([num, g, None,
                     min(wk["출발시각"]) if len(wk) else "평일표 없음",
                     max(wk["출발시각"]) if len(wk) else "평일표 없음",
                     max(hol["출발시각"]) if len(hol) else "공휴일표 없음(평일표만 게시)",
                     "공개 자료 없음 — 현2 실측", len(wk), len(hol),
                     f"{s['기점부터_정류장수'].min()}~{s['기점부터_정류장수'].max()}" if s['기점부터_정류장수'].nunique() > 1 else int(s['기점부터_정류장수'].iloc[0]),
                     f"{s['구간_정류장수'].min()}~{s['구간_정류장수'].max()}" if s['구간_정류장수'].nunique() > 1 else int(s['구간_정류장수'].iloc[0]),
                     " / ".join(sorted(set(s["시작정류장"]))) + " → " + " / ".join(sorted(set(s["끝정류장"]))),
                     len(names), "; ".join(names), BIS_TT, "누리집 게시(게시일 미표기)", TODAY, ""])
    order = {p: i for i, p in enumerate([f"{a}→{b}" for a, b in PAIRS])}
    rows.sort(key=lambda r: (order[r[1]], str(r[0])))
    header = ["노선번호", "구간", "시간대별 배차간격", "첫차", "막차(평일)", "막차(주말)", "소요시간",
              "운행 횟수(평일)", "운행 횟수(공휴일)", "기점→구간 시작 정류장 수", "구간 정류장 수", "구간 정류장(시작→끝)",
              "노선 변형 수", "노선명(변형)", "출처 URL", "게시일", "조사일", "비고"]
    idx = write_table(ws, header, rows, widths=[9, 14, 16, 8, 9, 22, 20, 9, 9, 11, 9, 34, 7, 60, 30, 20, 11, 30],
                      link_cols=["출처 URL"], date_cols=["조사일"], time_cols=["첫차", "막차(평일)", "막차(주말)"])
    c = {k: get_column_letter(v) for k, v in idx.items()}
    for r in range(2, ws.max_row + 1):
        # 평균 배차간격(분) = (막차−첫차)/(운행횟수−1). 여러 변형을 합친 기점 출발 기준이라 '평균'일 뿐
        ws.cell(r, idx["시간대별 배차간격"]).value = (
            f'=IF(AND(ISNUMBER({c["첫차"]}{r}),ISNUMBER({c["막차(평일)"]}{r}),{c["운행 횟수(평일)"]}{r}>1),'
            f'"평균 "&ROUND(({c["막차(평일)"]}{r}-{c["첫차"]}{r})*1440/({c["운행 횟수(평일)"]}{r}-1),0)&"분(시간대별은 옆 시트)","1회 이하")')
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, idx["노선번호"]).value == "112":
            ws.cell(r, idx["비고"]).value = "월영교 경유 유일 노선(2022·2023 노선집). 2023.10 노선집 전체 운행시간 50분·배차 65분"

    # 출발시각_전체 + 구간 플래그
    ds = wb.create_sheet("출발시각_전체")
    seg_names = [f"{a}→{b}" for a, b in PAIRS]
    drows = []
    for _, r in dep.sort_values(["노선번호", "routeId", "시간표 구분", "출발시각"]).iterrows():
        f = seg_by_rid.get(r.routeId, set())
        drows.append([r["노선번호"], r["노선명"], r["시간표 구분"], r["기점"], r["종점"], r["출발시각"], None,
                      *["Y" if s in f else "" for s in seg_names], r.routeId])
    dh = ["노선번호", "노선명", "시간표 구분", "기점", "종점", "출발시각", "시(時)", *seg_names, "routeId"]
    didx = write_table(ds, dh, drows, widths=[8, 50, 12, 22, 22, 8, 6] + [11] * 6 + [12], time_cols=["출발시각"])
    for r in range(2, ds.max_row + 1):
        ds.cell(r, didx["시(時)"]).value = f"=HOUR(F{r})"
    n_dep = ds.max_row

    # 시간대별_운행(평일) — COUNTIFS
    hs = wb.create_sheet("시간대별_운행(평일)")
    hours = list(range(5, 24))
    hs.append(["구간 \\ 기점 출발 시(時)"] + [f"{h:02d}시" for h in hours] + ["합계", "18시 이후", "21시 이후"])
    S = "'출발시각_전체'"
    hcol = get_column_letter(didx["시(時)"])
    tcol = get_column_letter(didx["시간표 구분"])
    for s in seg_names:
        r = hs.max_row + 1
        fcol = get_column_letter(didx[s])
        cells = [f'=COUNTIFS({S}!${fcol}$2:${fcol}${n_dep},"Y",{S}!${hcol}$2:${hcol}${n_dep},{h},{S}!${tcol}$2:${tcol}${n_dep},"평일")' for h in hours]
        last = get_column_letter(1 + len(hours))
        hs.append([s] + cells + [f"=SUM(B{r}:{last}{r})",
                                 f"=SUM({get_column_letter(2 + hours.index(18))}{r}:{last}{r})",
                                 f"=SUM({get_column_letter(2 + hours.index(21))}{r}:{last}{r})"])
    hs.append([])
    hs.append(["읽는 법", "기점 출발 시각 기준 평일 운행 횟수. 구간 시작 정류장 통과 시각은 기점부터 정류장 수만큼 늦다(누리집에 정류장별 시각 없음). 여러 노선 변형 합계."])
    for row in hs.iter_rows():
        for cc in row:
            cc.font = FONT
    for cc in hs[1]:
        cc.font = BOLD
    hs.column_dimensions["A"].width = 22

    # 112번 정류장 순서
    ps = wb.create_sheet("112번_정류장순서")
    ids112 = set(rt[rt.routeNum == "112"].routeId)
    lab = lambda sid: "/".join(nm for nm, s in area.items() if sid in s)
    prow = []
    for rid in sorted(ids112):
        g = rs[rs.routeId == rid].sort_values("routeOrd")
        nm = rt.loc[rt.routeId == rid, "routeNm"].iloc[0]
        times = dep[(dep.routeId == rid)]
        tstr = ", ".join(f"{t['시간표 구분']} {t['출발시각'].strftime('%H:%M')}" for _, t in times.iterrows()) or "시간표 없음"
        for _, s in g.iterrows():
            prow.append([nm, int(s.routeOrd), s.stopName, lab(s.stopId), s.stopId, tstr, rid])
    write_table(ps, ["노선명(변형)", "순서", "정류장명", "권역(기준점 반경 안)", "BIS stopId", "기점 출발 시각", "routeId"], prow,
                widths=[44, 6, 24, 12, 12, 40, 12])

    notes_sheet(wb, "출처·메모", [
        ("출처", f"안동시 버스정보시스템 누리집 '버스노선 시간표'({BIS_TT})가 불러오는 공개 JSON: /api/hp/busTimetable(시간표 405행·46개 노선), /api/hp/busRoute(노선 변형 422개), /api/bus/routes/stops(노선별 정류장 순서 20,642행). 2026-09-22 조회. 관리자·로그인 API는 쓰지 않음."),
        ("시간표 구분", "누리집 표기 그대로: 평일(holidayType 1) / 공휴일[일요일](holidayType 2). 토요일이 어느 표를 따르는지는 누리집에 없음 → 막차(주말)는 공휴일표 기준, 공휴일표가 없는 노선은 '공휴일표 없음(평일표만 게시)'으로 적고 운행 여부를 단정하지 않음."),
        ("구간 정의", "기준점(데이터랩 POI 좌표) 반경 안 정류장: 안동역 300m(안동역(안동터미널)·노하동입구 5개), 원도심 = 안동구시장 500m(교보생명·구시장입구·신시장·웅부공원·안동시청 등 20개), 월영교 300m(월영교 2개). 노선 정류장 순서에서 A권역 정류장 뒤에 B권역 정류장이 오면 'A→B' 구간 운행으로 봄."),
        ("첫차·막차", "기점 출발 시각이다. 누리집에 정류장별 통과 시각이 없어 구간 시작 정류장 통과는 더 늦다('기점→구간 시작 정류장 수' 열 참고)."),
        ("소요시간", "공개 자료 없음(누리집·API에 정류장 간 시간 없음) → 작업 목록 현2(원도심→월영교 실제 이동시간) 실측으로 채움. 참고: 2023.10 노선안내책자 112번 전체 운행시간 50분, 운행거리 20.0km."),
        ("핵심", "원도심↔월영교, 안동역↔월영교는 112번뿐. 이름에 월영교가 들어간 변형 '112(교보생명-월영교-문화관광단지)'는 평일 17:55·18:45 두 회. 2022.3 노선집 112번 막차 18:45(교보생명)/19:00과 같음."),
        ("TAGO", "국토교통부 TAGO 버스노선·정류소 API는 보유 키 미등록(403). 이번엔 안동시 BIS 누리집 공개 자료로 대체."),
        ("조사일", "2026-09-22"),
    ])
    path = OUT / "01_시내버스_노선시간표.xlsx"
    wb.save(path)
    return path, seg, dep


def build_04():
    stops, rs, tt, rt, area = load()
    # 국토부 파일과 결합(정류장번호 = 'ADB' + BIS stopId)
    molit_path = ROOT / "data/external/작업목록_1-12_보강_20260922/04_정류장/국토교통부_전국 버스정류장 위치정보_20251031.csv"
    m = pd.read_csv(molit_path, encoding="cp949", dtype=str)
    m = m[m["도시명"].str.contains("안동", na=False)]
    m_ids = set(m["정류장번호"])
    routes_at = rs.merge(rt[["routeId", "routeNum"]], on="routeId").groupby("stopId")["routeNum"].apply(lambda s: sorted(set(s))).to_dict()
    team = {"ADB354000049": "안동역", "ADB354000063": "안동역", "ADB354000416": "안동역", "ADB354000459": "안동역",
            "ADB354000536": "안동역", "ADB354000716": "월영교", "ADB354000843": "월영교", "ADB354000943": "원도심(문화의거리)",
            "ADB354001806": "원도심(문화의거리)", "ADB354000965": "원도심(음식의거리 등)"}
    rows = []
    for nm, (x, y, rad) in {"안동역": (128.6748, 36.5747, 300), "원도심": (128.728, 36.5655, 500), "월영교": (128.7609, 36.5765, 500)}.items():
        sel = stops[stops["d_" + nm] <= rad].sort_values("d_" + nm) if "d_" + nm in stops else None
        if sel is None:
            d = hav(x, y, stops.lng.values, stops.lat.values)
            sel = stops[d <= rad].assign(**{"d_" + nm: d[d <= rad]}).sort_values("d_" + nm)
        for _, s in sel.iterrows():
            code = f"ADB{s.stopId}"
            r_list = routes_at.get(s.stopId, [])
            tnote = team.get(code, "")
            note = ""
            if tnote and not tnote.startswith(nm):
                note = f"팀 파일에는 '{tnote}'로 분류 — 좌표상 {nm}"
            rows.append([s.stopName, s.lat, s.lng, "설치" if s.bitYn == "Y" else "미설치", nm, round(float(s["d_" + nm])),
                         code, int(s.stopId), s.serviceId, len(r_list), ", ".join(r_list),
                         "예" if code in team else "", "있음" if code in m_ids else "없음",
                         BIS, "누리집 게시(게시일 미표기)", TODAY, note])
    # 팀 추출 중 반경 밖(구 안동역 2곳 등)
    have = {r[6] for r in rows}
    for code, lab_ in team.items():
        if code in have:
            continue
        s = stops[stops.stopId == int(code[3:])].iloc[0]
        d = {nm: hav(x, y, s.lng, s.lat) for nm, (x, y, _) in ANCHOR.items()}
        near = min(d, key=d.get)
        r_list = routes_at.get(s.stopId, [])
        rows.append([s.stopName, s.lat, s.lng, "설치" if s.bitYn == "Y" else "미설치", f"반경 밖(가장 가까운 기준점: {near})",
                     round(float(d[near])), code, int(s.stopId), s.serviceId, len(r_list), ", ".join(r_list), "예",
                     "있음" if code in m_ids else "없음", BIS, "누리집 게시(게시일 미표기)", TODAY,
                     f"팀 파일에는 '{lab_}'로 분류. 옛 안동역 자리(현재 안동역까지 약 5.4km, 구시장까지 약 600m)" if "구 안동역" in s.stopName else f"팀 파일 분류 '{lab_}'"])
    wb = Workbook()
    ws = wb.active
    ws.title = "권역별_정류장"
    header = ["정류장명", "위도", "경도", "안내기 설치 여부", "권역", "기준점까지 거리(m)", "정류장번호(국토부)", "BIS stopId",
              "모바일단축번호", "경유 노선 수", "경유 노선", "팀 추출", "국토부 파일", "출처 URL", "게시일", "조사일", "비고"]
    write_table(ws, header, rows, widths=[20, 10, 11, 10, 16, 9, 15, 11, 9, 7, 44, 6, 8, 26, 20, 11, 44],
                link_cols=["출처 URL"], date_cols=["조사일"])
    n = ws.max_row
    ws.append([])
    ws.append(["권역", "정류장 수", "안내기 설치"])
    for nm in ["안동역", "원도심", "월영교"]:
        r = ws.max_row + 1
        ws.append([nm, f'=COUNTIFS($E$2:$E${n},A{r})', f'=COUNTIFS($E$2:$E${n},A{r},$D$2:$D${n},"설치")'])
    for row in ws.iter_rows(min_row=n + 1):
        for c in row:
            c.font = FONT

    al = wb.create_sheet("안동_전체_정류장")
    arows = [[s.stopName, s.lat, s.lng, "설치" if s.bitYn == "Y" else "미설치", f"ADB{s.stopId}", int(s.stopId), s.serviceId,
              "있음" if f"ADB{s.stopId}" in m_ids else "없음", BIS, TODAY] for _, s in stops.sort_values("stopName").iterrows()]
    write_table(al, ["정류장명", "위도", "경도", "안내기 설치 여부", "정류장번호(국토부)", "BIS stopId", "모바일단축번호", "국토부 파일", "출처 URL", "조사일"],
                arows, widths=[24, 10, 11, 10, 15, 11, 9, 8, 26, 11], date_cols=["조사일"])
    k = al.max_row
    al.append([])
    al.append(["전체 정류장", f"=COUNTA(A2:A{k})"])
    al.append(["안내기 설치", f'=COUNTIFS(D2:D{k},"설치")'])
    al.append(["설치 비율", f"=B{k + 3}/B{k + 2}"])
    al.cell(al.max_row, 2).number_format = "0.0%"
    for row in al.iter_rows(min_row=k + 1):
        for c in row:
            c.font = FONT

    notes_sheet(wb, "출처·메모", [
        ("안내기 설치 여부", f"안동시 버스정보시스템 누리집 지도가 불러오는 공개 정류장 목록(/api/bus/stops)의 bitYn 값. 누리집은 bitYn=Y인 정류장에 BIT 아이콘을 그린다. {BIS} 2026-09-22 조회. 3,305개 중 316개 'Y'."),
        ("위치", f"BIS 좌표 사용. 국토교통부 전국 버스정류장 위치정보(2025-10-31, {MOLIT} — 팀이 받은 파일 ~/Downloads/국토교통부_전국 버스정류장 위치정보_안동시/)와 정류장번호로 대조('국토부 파일' 열). 국토부 파일의 '(미정차)' 경기BIS 27개는 안동 정차 정류장이 아니라 제외."),
        ("권역 반경", "안동역 300m · 원도심(안동구시장 기준) 500m · 월영교 500m. 기준점 = 데이터랩 POI 좌표. '방향'은 같은 이름의 길 건너편 정류장을 모두 포함하는 것으로 처리."),
        ("팀 추출 검증", "팀 파일 10개 중 '안동역' 파일의 '구 안동역'·'구 안동역 건너' 2곳은 옛 안동역 자리(원도심 쪽)라 안동역 권역이 아님. 원도심은 구시장 300m 안에만 11개, 500m 안 20개인데 팀 파일엔 3개."),
        ("경유 노선", "/api/bus/routes/stops(노선 변형 422개)의 정류장 순서에서 그 정류장을 지나는 노선 번호."),
        ("조사일", "2026-09-22"),
    ])
    path = OUT / "04_BIS정류장_위치.xlsx"
    wb.save(path)
    return path, rows


# ───────────────────────── 02 주민증: 시작 시기·발급 건수 보강(기존 파일 수정) ─────────────────────────
KTO_2306 = "https://knto.or.kr/pressRelease/547483"
EDAILY_2303 = "https://edaily.co.kr/News/Read?mediaCodeNo=257&newsId=01610486635541680"
SEOUL_2311 = "https://www.seoul.co.kr/news/society/2023/11/30/20231130500086"
JC_2406 = "https://jc-one.co.kr/jc-news/%EC%A0%9C%EC%B2%9C-%EB%94%94%EC%A7%80%ED%84%B8-%EA%B4%80%EA%B4%91%EC%A3%BC%EB%AF%BC%EC%A6%9D-%EB%B0%9C%EA%B8%89%EC%9E%90-%EC%88%98-5%EB%A7%8C%EB%AA%85-%EB%8F%8C%ED%8C%8C/"
KOREA_2406 = "https://www.korea.kr/multi/visualNewsView.do?newsId=148929998"
MCST_2503 = ("https://www.mcst.go.kr/servlets/eduport/front/upload/UplDownloadFile?pFileName=%280326%29%EB%AC%B8%EC%B2%B4%EB%B6%80%EB%B3%B4%EB%8F%84%EC%9E%90%EB%A3%8C-"
             "%EB%94%94%EC%A7%80%ED%84%B8_%EA%B4%80%EA%B4%91%EC%A3%BC%EB%AF%BC%EC%A6%9D_%EC%9A%B4%EC%98%81_%EC%A7%80%EC%97%AD_%ED%99%95%EB%8C%80.pdf"
             "&pRealName=20250326080436339791604117_PRESS20250326080436485717.pdf&pPath=0302000000")
NEWDAILY_2606 = "https://www.newdaily.co.kr/site/data/html/2026/06/08/2026060800045.html"
SEGYE_2409 = "https://jeju.thesegye.com/news/view/1065617526180329?dt=m"
KHAN_2310 = "https://www.khan.co.kr/article/202310231428001"

PHASES = [  # (시작 시기, 근거, 출처, 게시일, 지역)
    ("2022-10(시범)", "관광공사 보도자료: 평창·옥천 '2022년 10월부터' 시범", KTO_2306, "2023-06-01", ["평창군", "옥천군"]),
    ("2023-05-31", "관광공사 보도자료: '5월 31일(수)부터' 11개 지역으로 확대(신규 9곳 명단은 이데일리 2023-03-09)", KTO_2306, "2023-06-01",
     ["강화군", "정선군", "단양군", "태안군", "고창군", "신안군", "고령군", "거창군", "부산광역시 영도구"]),
    ("2023-10-25 추가 선정", "서울신문: 관광공사가 '지난달 25일' 4개 지역 추가 선정 → 15개 지역. 운영 시작 월은 제천만 확인('11월부터', 제천또바기뉴스 2024-06-15)", SEOUL_2311, "2023-11-30",
     ["연천군", "제천시", "남원시", "하동군"]),
    ("2024-06", "정책브리핑 카드뉴스: 15개 지역 → 34개 지역(6월). 신규 19곳 명단은 공식 명단 두 개의 차이(2023년 15곳 vs 2025-03-26 문체부 44곳 중 2025 신규 10곳 제외)로 도출", KOREA_2406, "2024-06-07",
     ["삼척시", "양양군", "영월군", "태백시", "홍천군", "괴산군", "영동군", "예산군", "가평군", "무주군", "임실군", "영광군", "장흥군", "해남군",
      "안동시", "영덕군", "영주시", "합천군", "부산광역시 서구"]),
    ("2025-04 말", "문체부 보도자료: 4월 말부터 34곳 → 44곳, 신규 10곳 명단", MCST_2503, "2025-03-26",
     ["철원군", "보령시", "김제시", "구례군", "곡성군", "함평군", "청도군", "의성군", "밀양시", "부산광역시 동구"]),
    ("2026-06-08", "뉴데일리: 신규 8곳 '이달 8일부터' 서비스 개시 → 52곳", NEWDAILY_2606, "2026-06-08",
     ["보은군", "순창군", "고흥군", "담양군", "완도군", "울진군", "산청군", "함양군"]),
]
ISSUE = {  # 지자체별 발급 수치(공개 보도에서 확인된 것만)
    "옥천군": ("47,861명(2023-10-18 기준)", KHAN_2310),
    "단양군": ("25,000명(2023-10-11 기준)", KHAN_2310),
    "제천시": ("5만 명 돌파(2024-06-12 기준)", JC_2406),
    "안동시": ("6만1천여 명(2024-09-21 보도 시점)", SEGYE_2409),
    "평창군": ("개별 수치 못 찾음 — 평창·옥천 합계 58,000여 명(2022.10~2023.4)", KTO_2306),
}


def build_02():
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment
    path = OUT / "02_디지털관광주민증_운영지자체.xlsx"
    wb = load_workbook(path)
    ws = wb["운영지자체"]
    hdr = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(hdr)}
    # 이미 보강한 파일이면 새 열을 다시 만들지 않는다
    for h in ["시작 시기 근거", "시작 시기 출처 URL", "시작 시기 게시일", "발급 건수 출처 URL"]:
        if h not in col:
            c = ws.cell(1, ws.max_column + 1, h)
            c.font, c.fill, c.alignment = base.BOLD, base.HEAD_FILL, Alignment(wrap_text=True, vertical="center")
            col[h] = c.column
    phase_of = {nm: p for p in PHASES for nm in p[4]}
    done, missing = 0, []
    for r in range(2, ws.max_row + 1):
        nm = ws.cell(r, col["지자체명"]).value
        if not nm or nm.startswith(("합계", "지자체 수")):
            continue
        p = phase_of.get(nm)
        if not p:
            missing.append(nm)
            continue
        start = "2024-06-01" if nm == "안동시" else ("2023-11" if nm == "제천시" else p[0])
        ws.cell(r, col["시작 시기"]).value = start
        ws.cell(r, col["시작 시기 근거"]).value = p[1] + (" · 안동은 '6월 1일부터'(제주세계타임즈 2024-09-21)" if nm == "안동시" else "")
        ws.cell(r, col["시작 시기 출처 URL"]).value = p[2]
        ws.cell(r, col["시작 시기 게시일"]).value = p[3]
        if nm in ISSUE:
            ws.cell(r, col["발급 건수"]).value = ISSUE[nm][0]
            ws.cell(r, col["발급 건수 출처 URL"]).value = ISSUE[nm][1]
        else:
            ws.cell(r, col["발급 건수"]).value = "못 찾음 — 지역별 공개 수치 없음(전국 누적만 공개, '출처·메모' 참고)"
        ws.cell(r, col["조사일"]).value = TODAY
        ws.cell(r, col["조사일"]).number_format = "yyyy-mm-dd"
        for h in ["시작 시기", "발급 건수", "시작 시기 근거", "시작 시기 출처 URL", "시작 시기 게시일", "발급 건수 출처 URL"]:
            c = ws.cell(r, col[h])
            c.font, c.alignment = base.FONT, base.WRAP
            if isinstance(c.value, str) and c.value.startswith("http"):
                c.hyperlink, c.font = c.value, base.LINK
        done += 1
    for h, w in [("시작 시기", 16), ("발급 건수", 30), ("시작 시기 근거", 50), ("시작 시기 출처 URL", 30), ("시작 시기 게시일", 11), ("발급 건수 출처 URL", 30)]:
        ws.column_dimensions[get_column_letter(col[h])].width = w
    assert not missing, missing

    # 확대 연혁 시트
    if "확대 연혁" in wb.sheetnames:
        del wb["확대 연혁"]
    hs = wb.create_sheet("확대 연혁", index=1)
    cum, rows = 0, []
    for start, why, url, pub, names in PHASES:
        cum += len(names)
        rows.append([start, len(names), cum, ", ".join(names), why, url, pub, TODAY])
    write_table(hs, ["시작 시기", "신규 지역 수", "누적 지역 수", "신규 지역", "근거", "출처 URL", "게시일", "조사일"], rows,
                widths=[16, 9, 9, 60, 60, 36, 11, 11], link_cols=["출처 URL"], date_cols=["조사일"])
    n = hs.max_row
    hs.append([])
    hs.append(["검증", f"=SUM(B2:B{n})", "← 52이어야 함(운영지자체 시트 지자체 수)"])
    hs.append(["전국 누적 발급", "411만 건(2022~2024-12 말)", "문체부 2025-03-26 보도자료", MCST_2503])
    hs.append(["11개 지역 발급·이용", "발급 325,165명, 실제 이용 29,785명(9.2%) — 2023-10 말", "서울신문 2023-11-30", SEOUL_2311])
    for row in hs.iter_rows(min_row=n + 1):
        for c in row:
            c.font, c.alignment = base.FONT, base.WRAP

    m = wb["출처·메모"]
    for r in range(2, m.max_row + 1):
        k = m.cell(r, 1).value
        if k == "시작 시기":
            m.cell(r, 2).value = ("(9/21) 공식 지역 페이지에 개시일 필드 없음 → (9/22 보강) 확대 발표 6건으로 52곳 전부 채움: 2022-10 시범 2 · 2023-05-31 9 · "
                                  "2023-10-25 선정 4 · 2024-06 19 · 2025-04 말 10 · 2026-06-08 8. '확대 연혁' 시트.")
        if k == "발급 건수":
            m.cell(r, 2).value = ("비로그인 공식 화면에 없음. (9/22 보강) 보도로 확인된 것만: 옥천 47,861명(2023-10-18)·단양 25,000명(2023-10-11)·제천 5만 명(2024-06-12)·"
                                  "안동 6만1천여 명(2024-09-21 보도). 전국 누적 411만 건(2024-12 말). 최신 안동 발급 수는 정보공개청구 14번 항목.")
        if k == "조사일":
            m.cell(r, 2).value = "2026-09-21(목록·혜택업체) / 2026-09-22(시작 시기·발급 건수 보강)"
    wb.save(path)
    return path, done


# ───────────────────────── 03 체험 프로그램: 가격·운영시간 보강(기존 파일 수정) ─────────────────────────
T_372 = "https://www.tourandong.com/public/sub1/sub1_5.cshtml?seq=372"
T_1545 = "https://www.tourandong.com/public/sub2/sub3.cshtml?seq=1545&page=1&searchKey=0&search="
SOJU_MALL = "https://andongsojumall.com/brand/experience.html"
HAHOE_RSV = "http://www.hahoemask.co.kr/bbs/board.php?bo_table=reserv&mode=step1"
HAHOE_SHOW = "http://www.hahoemask.co.kr/bbs/content.php?co_id=camp"
MOON_OP = "http://koreaboat.co.kr/htm/menu2_01.htm"
MOON_ADDMO = "https://addmo.kr/public/sub5/sub1_1.asp?id=1050"
MOON_TELL = "https://www.telltrip.com/domestic-travel/andong-wolyeonggyo-summer-night-walk/"
GRAND_MASK = "https://andong.grandculture.net/andong/toc/GC02430063"


def build_03():
    from openpyxl import load_workbook
    path = OUT / "03_체험프로그램_목록가격.xlsx"
    wb = load_workbook(path)
    ws = wb["체험프로그램"]
    hdr = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(hdr)}
    rows = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value}

    def put(name_prefix, **kw):
        r = next(v for k, v in rows.items() if k.startswith(name_prefix))
        for h, val in kw.items():
            if h == "비고+":
                old = ws.cell(r, col["비고"]).value or ""
                if val not in old:
                    ws.cell(r, col["비고"]).value = (old + " / " if old else "") + val
            else:
                ws.cell(r, col[h]).value = val
        ws.cell(r, col["조사일"]).value = TODAY
        ws.cell(r, col["조사일"]).number_format = "yyyy-mm-dd"

    put("하회세계탈박물관 관람", **{"비고+": f"(9/22) 하회동탈박물관 = 같은 시설 — 디지털안동문화대전 '하회동 탈 박물관' 이칭 '하회세계탈박물관'({GRAND_MASK}). 데이터랩 주요관광지 통계에는 '하회동탈박물관'으로 올라 있음"})
    put("하회별신굿탈놀이 상설공연", **{"비고+": f"(9/22) 보존회 누리집: 3~12월 화~일·1~2월 토·일 14:00~15:00(1시간) {HAHOE_SHOW}"})
    put("탈춤따라배우기", **{"가격": "보존회 예약 페이지: 탈춤 아카데미·탈 만들기·탈춤 따라배우기는 30명 기준 기본 30만원(초과 인원 추가). 1박2일 체험 캠프 50,000원",
                         "예약 방법": f"보존회 누리집 온라인 예약(달력 선택) {HAHOE_RSV}",
                         "비고+": "(9/22) 30명 단위 단체 프로그램 — 개인 저녁 연결용이 아님"})
    put("명인 안동소주 양조장체험", **{"가격": "칵테일 15,000원(약 1시간 30분) · 누룩 만들기 20,000원 · 전통주 빚기 25,000원 · 곡류 발효음료 25,000원(각 약 2시간, 1인) · 시음·전시관 무료(20분 이내)",
                              "예약 방법": "전화 신청 후 결제(054-856-6903). 체험 인원 15명 이상 단체",
                              "비고+": f"(9/22) 운영사 누리집 {SOJU_MALL} — 운영 요일·시간은 미게시. 15명 이상 단체 조건이라 개인 방문객 연결 불가"})
    put("민속주 안동소주 만들기 체험", **{"비고+": "(9/22) 안동관광 체험 페이지·경향신문 2025-10-06에도 체험비 없음 → 054-858-4541 문의 필요"})
    put("월영교 문보트 체험", **{"가격": f"공식 요금 미게시(운영사·안동시·관광협의회). 여행 매체 기재 '3인 기준 28,000원/30분'({MOON_TELL}) — 확인 필요. 10인 이상 10% 할인(관광협의회)",
                            "휴무": "월요일(운영사·관광협의회)",
                            "비고+": f"(9/22) 운영시간은 계절별로 다름: 운영사 (주)글로벌코리아 게시(2026.1.1~2.28 기준) 화~금 10:00~22:00, 토·일·공휴일 11:00~23:00 {MOON_OP} / 안동시관광협의회 '안동을 담다' 주중 10:00~23:00·주말 10:00~24:00 {MOON_ADDMO}. 문의 054-853-0715"})

    new = [
        ["안동공예문화전시관 공예체험(도자기·한지·천연염색·가죽·우드아트 등)", "안동공예문화전시관", "안동시 석주로 245(월영교 인근)",
         "도자기 10,000~20,000원 · 한지 5,000원~ · 천연염색 8,000~25,000원 · 보석십자수/폼블럭 10,000원~ · 천아트 10,000~39,000원 · 가죽 12,000~23,000원 · 아트페인팅 8,000~15,000원 · 우드아트 6,000~15,000원 · 금속 12,000원~",
         "명시 없음", "09:00", "18:00", "월요일·1월 1일·설·추석 연휴(전후 3일)·선거일", "개인 가능. 단체(10명 이상)는 일주일 전 연락, 30명 이상 20% 할인. 문의 054-843-5531",
         "N(18:00 종료)", T_372, "페이지 게시일 미표기", TODAY, "(9/22 추가) 월영교에서 가장 가까운 공예 체험. 관람 무료"],
        ["전통리조트 구름에 체험(고추장·두부·국시·명인 하이볼)", "전통리조트 구름에 On", "안동시 민속촌길 190(월영교 인근)",
         "수제 고추장·우리콩 두부·안동 국시 각 22,000원/1인 · 명인 하이볼 12,000원/1인", "1개당 1시간 내외", "명시 없음", "명시 없음", "연중무휴",
         "전화 문의 054-823-9001", "확인 필요(운영 종료 시각 미게시)", T_1545, "페이지 게시일 미표기", TODAY,
         "(9/22 추가) 월영교 옆 숙박시설 체험. 저녁 운영 여부는 전화로 확인해야 1단계 저녁 연결에 쓸 수 있음"],
    ]
    have = set(rows)
    for r in new:
        if r[0] in have:
            continue
        ws.append(r)
        rr = ws.max_row
        for c in ws[rr]:
            c.font, c.alignment = base.FONT, base.WRAP
        u = ws.cell(rr, col["출처 URL"])
        u.hyperlink, u.font = u.value, base.LINK
        ws.cell(rr, col["조사일"]).number_format = "yyyy-mm-dd"
    for r in range(2, ws.max_row + 1):
        for h in ("가격", "휴무", "예약 방법", "비고"):
            c = ws.cell(r, col[h])
            c.font, c.alignment = base.FONT, base.WRAP

    m = wb["출처·메모"]
    notes = {k: r for r in range(2, m.max_row + 1) for k in [m.cell(r, 1).value]}
    line = ("(9/22 보강) 명인 안동소주 가격(단체 15명 이상)·보존회 탈 체험 가격(30명 기준 30만원)·문보트 운영시간 3종 대조·하회동탈박물관=하회세계탈박물관 확인, "
            "월영교 인근 체험 2곳(안동공예문화전시관 09~18시, 구름에 종료시각 미게시) 추가. 여전히 못 찾음: 안동소주박물관 체험비, 선성현 의복체험 가격·시간, 문보트 공식 요금, 안동호반힐링타운 가격.")
    if "9/22 보강" not in notes:
        m.append(["9/22 보강", line])
        for c in m[m.max_row]:
            c.font, c.alignment = base.FONT, base.WRAP
    if "핵심 판정" in notes:
        r = notes["핵심 판정"]
        v = m.cell(r, 2).value or ""
        if "(9/22)" not in v:
            m.cell(r, 2).value = v + " (9/22) 추가 확인 후에도 18시 이후 종료가 확인된 체험은 월영교 문보트·황포돛배뿐. 구름에 체험은 종료 시각 미게시."
    wb.save(path)
    return path


# ───────────────────────── 05 주요관광지점 입장객(유료/무료) ─────────────────────────
KCTI = "https://know.tour.go.kr/stat/visitStatDis/table.do"
RAW5 = ROOT / "data/external/작업목록_1-12_보강_20260922/05_입장객"


def parse_kcti(fn):
    import re
    s = open(fn, encoding="utf-8").read()
    out = []
    for r in re.findall(r"<row[^>]*>(.*?)</row>", s, flags=re.S):
        d = dict(re.findall(r"<([A-Z_0-9]+)>(.*?)</\1>", r, flags=re.S))
        nm = d.get("RES_NM", "")
        code = re.search(r"detail\('([^']+)'\)", nm)
        d["CODE"] = code.group(1) if code else ""
        d["NAME"] = re.sub(r"<[^>]+>", "", re.sub(r"<!\[CDATA\[|\]\]>", "", nm)).strip()
        out.append(d)
    return out


def build_05():
    long_rows = {}
    for tag, yrs in [("", "2023-2026"), ("2018-2022_", "2018-2022")]:
        A = parse_kcti(RAW5 / f"statTableData_{tag}A.xml")
        paid = {d["CODE"] for d in parse_kcti(RAW5 / f"statTableData_{tag}C.xml")}
        free = {d["CODE"] for d in parse_kcti(RAW5 / f"statTableData_{tag}F.xml")}
        rows = []
        by = {}
        for d in A:
            by.setdefault((d["CODE"], d["NAME"]), {})[d["NF_GB"]] = d
        for (code, name), g in sorted(by.items(), key=lambda kv: kv[0][1]):
            cls = "유료" if code in paid else ("무료" if code in free else "미분류")
            tot = g.get("합계") or g.get("내국인")
            months = [k for k in tot if k.startswith("M_")]
            for k in months:
                ym = f"{k[2:6]}-{k[6:8]}"
                v = lambda gb: (None if (gb not in g or g[gb][k] in ("-", "")) else int(g[gb][k]))
                t, dom, fr = v("합계"), v("내국인"), v("외국인")
                if t is None and dom is not None:
                    t = dom + (fr or 0)
                rows.append([name, ym, t if cls == "유료" else None, t if cls == "무료" else None, t, dom, fr, cls, code,
                             KCTI, "자료 갱신일 2026-08-18(2026년 2분기 잠정치)", TODAY, "" if t is not None else "미보고('-')"])
        long_rows[yrs] = rows

    wb = Workbook()
    ws = wb.active
    ws.title = "지점별_월별(2023-2026.6)"
    header = ["지점명", "연월", "유료 입장객", "무료 입장객", "합계", "내국인", "외국인", "유료/무료 구분", "지점코드", "출처 URL", "게시일", "조사일", "비고"]
    fmt = {k: "#,##0" for k in ["유료 입장객", "무료 입장객", "합계", "내국인", "외국인"]}
    write_table(ws, header, long_rows["2023-2026"], widths=[24, 9, 11, 11, 11, 11, 9, 9, 14, 30, 24, 11, 14],
                link_cols=["출처 URL"], date_cols=["조사일"], num_cols=fmt)
    n1 = ws.max_row
    ws2 = wb.create_sheet("지점별_월별(2018-2022 참고)")
    write_table(ws2, header, long_rows["2018-2022"], widths=[24, 9, 11, 11, 11, 11, 9, 9, 14, 30, 24, 11, 14],
                link_cols=["출처 URL"], date_cols=["조사일"], num_cols=fmt)
    n2 = ws2.max_row

    # 연도별 요약(SUMIFS) — 보고 개월 수를 같이 보여 부분 연도 비교를 막는다
    sm = wb.create_sheet("연도별_요약", index=0)
    names = sorted({r[0] for r in long_rows["2023-2026"]} | {r[0] for r in long_rows["2018-2022"]})
    cls_of = {}
    for yrs in ("2018-2022", "2023-2026"):
        for r in long_rows[yrs]:
            cls_of[r[0]] = r[7]
    years = [2018, 2019, 2020, 2021, 2022, 2023, 2024, 2025, 2026]
    head = ["지점명", "유료/무료(최근)"] + [f"{y}{' (1~6월)' if y == 2026 else ''} 합계" for y in years] + [f"{y} 보고 개월" for y in years]
    sm.append(head)
    S1, S2 = "'지점별_월별(2023-2026.6)'", "'지점별_월별(2018-2022 참고)'"
    for nm in names:
        r = sm.max_row + 1
        vals = [nm, cls_of.get(nm, "")]
        for y in years:
            S, n = (S2, n2) if y <= 2022 else (S1, n1)
            vals.append(f'=SUMIFS({S}!$E$2:$E${n},{S}!$A$2:$A${n},$A{r},{S}!$B$2:$B${n},"{y}-*")')
        for y in years:
            S, n = (S2, n2) if y <= 2022 else (S1, n1)
            vals.append(f'=COUNTIFS({S}!$A$2:$A${n},$A{r},{S}!$B$2:$B${n},"{y}-*",{S}!$E$2:$E${n},"<>")')
        sm.append(vals)
    for row in sm.iter_rows(min_row=2):
        for c in row[2:2 + len(years)]:
            c.number_format = "#,##0"
    for row in sm.iter_rows():
        for c in row:
            c.font = FONT
    for c in sm[1]:
        c.font, c.fill = BOLD, base.HEAD_FILL
    sm.column_dimensions["A"].width = 26
    sm.freeze_panes = "C2"

    # 데이터랩 대조(보유 LN_03_012_001_001)
    dl = pd.read_csv(ROOT / "data/datalab_추가/주요관광지_월별방문자_전체_LN_03_012_001_001.csv", dtype={"BASE_YM": str})
    dl = dl[dl.SGG_NM == "안동시"]
    dl["ALL"] = pd.to_numeric(dl["ALL_TOU_NUM"], errors="coerce")
    kc = pd.DataFrame(long_rows["2023-2026"], columns=header)
    kc["연"] = kc["연월"].str[:4]
    cmp_rows = []
    for nm in sorted(set(dl.TAR_ID) & set(kc["지점명"])):
        for y in ("2023", "2024", "2025"):
            a = kc[(kc["지점명"] == nm) & (kc["연"] == y)]["합계"].dropna()
            b = dl[(dl.TAR_ID == nm) & (dl.BASE_YM.str.startswith(y))]["ALL"].dropna()
            if len(a) or len(b):
                cmp_rows.append([nm, y, int(a.sum()) if len(a) else None, len(a), int(b.sum()) if len(b) else None, len(b), None])
    cp = wb.create_sheet("데이터랩_대조")
    write_table(cp, ["지점명", "연도", "관광지식정보시스템 합계", "보고 개월", "데이터랩 LN_03_012 합계", "값 있는 개월", "일치"], cmp_rows,
                widths=[24, 7, 16, 9, 16, 10, 8], num_cols={"관광지식정보시스템 합계": "#,##0", "데이터랩 LN_03_012 합계": "#,##0"})
    for r in range(2, cp.max_row + 1):
        cp.cell(r, 7).value = f'=IF(OR(C{r}="",E{r}=""),"",IF(C{r}=E{r},"일치","다름"))'
    only_kcti = sorted(set(kc["지점명"]) - set(dl.TAR_ID))
    notes_sheet(wb, "출처·메모", [
        ("출처", f"문화체육관광부·한국문화관광연구원 관광지식정보시스템 「주요관광지점 입장객 통계」 통계표({KCTI}) — 경상북도 안동시, 월별, 분류 전체/유료/무료로 각각 조회(화면이 부르는 statTableData.do 원자료 XML 저장). 자료 갱신일 2026-08-18, 2026년 2분기 잠정치까지. 통계청 승인통계 제113005호."),
        ("유료/무료", "지점 단위 분류(입장료 있는 관광지 = 유료)다. 한 지점 안에서 유료·무료 방문을 나눈 값이 아니다. 그래서 유료 지점 행은 '유료 입장객'에, 무료 지점 행은 '무료 입장객'에 합계를 넣었다. 2023~2026 기준 38개 지점: 유료 20·무료 18."),
        ("미보고", "'-'는 그 달 보고가 없다는 뜻(0이 아님). '연도별_요약'의 '보고 개월'이 12 미만이면 연간 비교에 쓰지 않는다."),
        ("월영교(새 발견)", "월영교가 무료 지점으로 올라 있다: 2022년 388,156명·2023년 682,541명(두 해 모두 12개월 보고), 2018~2021년과 2024년 이후는 미보고. → 프로젝트 기록의 '월영교 방문자 수 원자료 없음'은 '2022~2023년 두 해만 있음'으로 고쳐야 한다. 무료 개방 공간이라 집계 방식은 통계표에 없음(지자체 보고) — 추세 근거로는 쓰지 말고 규모 참고로만."),
        ("데이터랩과의 관계", "데이터랩 LN_03_012(보유)는 29개 지점, 관광지식정보시스템은 38개 지점. 데이터랩에 없는 지점: " + ", ".join(only_kcti) + ". 겹치는 지점은 '데이터랩_대조' 시트."),
        ("주의", "통계표 안내 문구: 개별 관광지 입장객 집계이며 지자체 관광객 총량으로 쓸 수 없다."),
        ("조사일", "2026-09-22"),
    ])
    path = OUT / "05_주요관광지점_입장객.xlsx"
    wb.save(path)
    return path, long_rows


if __name__ == "__main__":
    import sys
    which = sys.argv[1:] or ["01", "04", "02", "03", "05"]
    if "05" in which:
        p5, lr = build_05()
        print(p5.name, {k: len(v) for k, v in lr.items()})
    if "03" in which:
        print(build_03().name, "보강")
    if "01" in which:
        p1, seg, dep = build_01()
        print(p1.name, "구간×변형", len(seg), "출발", len(dep))
    if "04" in which:
        p4, rows4 = build_04()
        print(p4.name, len(rows4), "rows")
    if "02" in which:
        p2, n2 = build_02()
        print(p2.name, "보강", n2, "곳")
